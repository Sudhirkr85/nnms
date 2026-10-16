import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = r"d:\Work\E-Learning"
BASE_SCIENCE_DIR = r"C:\Users\Orange\Downloads\New folder\Science_Complete_Course_SagarSir"
PHYSICS_DIR = os.path.join(BASE_SCIENCE_DIR, "01_Physics_भौतिक_विज्ञान")
CHEMISTRY_DIR = os.path.join(BASE_SCIENCE_DIR, "02_Chemistry_रसायन_विज्ञान")
BIOLOGY_DIR = os.path.join(BASE_SCIENCE_DIR, "03_Biology_जीव_विज्ञान")
MAHAMANTHAN_DIR = os.path.join(BASE_SCIENCE_DIR, "04_Science_MahaManthan_मॉक_टेस्ट्स")

for d in [PHYSICS_DIR, CHEMISTRY_DIR, BIOLOGY_DIR, MAHAMANTHAN_DIR]:
    os.makedirs(d, exist_ok=True)

LOGO_PATH = os.path.join(ROOT_DIR, "public", "logo-circle-transparent.png")
if not os.path.exists(LOGO_PATH):
    LOGO_PATH = r"C:\Users\Orange\Downloads\New folder\assets\logo-circle-transparent.png"

# Science Color Palette
C_NAVY = RGBColor(15, 23, 42)          # Deep Slate/Navy Primary
C_TEAL = RGBColor(13, 148, 136)        # Science Teal Accent
C_GOLD = RGBColor(217, 119, 6)         # Amber Accent
C_LIGHT_BG = RGBColor(248, 250, 252)   # Soft Light BG
C_WHITE = RGBColor(255, 255, 255)
C_DARK = RGBColor(30, 41, 59)          # Dark Body Text
C_MUTED = RGBColor(100, 116, 139)      # Muted Text
C_BLUE_CARD = RGBColor(240, 249, 255)  # Sky Light Card
C_AMBER_CARD = RGBColor(254, 243, 199) # Amber Card
C_GREEN = RGBColor(22, 101, 52)        # Correct Option Green Glow (#166534)
C_RED = RGBColor(185, 28, 28)

def create_base_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_header_footer(slide, title_text, slide_num, total_slides):
    # Header Background Card
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.2), Inches(12.533), Inches(0.85)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = C_NAVY
    header.line.color.rgb = C_TEAL
    header.line.width = Pt(1.5)

    # Logo in header (0.55 in to 1.30 in)
    if os.path.exists(LOGO_PATH):
        try:
            slide.shapes.add_picture(LOGO_PATH, Inches(0.55), Inches(0.25), Inches(0.75), Inches(0.75))
        except Exception:
            pass

    # Header Text in separate textbox (starts at 1.45 in, no overlap!)
    tb = slide.shapes.add_textbox(Inches(1.45), Inches(0.2), Inches(11.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)

    p0 = tf.paragraphs[0]
    p0.text = "SAGAR COACHING CENTRE BHAGWANPUR • NMMS All India (विज्ञान Special)"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = C_GOLD

    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE

    # Footer
    footer_text = f"WhatsApp / Helpline: +91 91101 13671  |  वेबसाइट: sagarcoaching.tech  |  Slide {slide_num} of {total_slides}"
    ft_box = slide.shapes.add_textbox(Inches(0.4), Inches(7.05), Inches(12.533), Inches(0.35))
    ft_tf = ft_box.text_frame
    ft_p = ft_tf.paragraphs[0]
    ft_p.text = footer_text
    ft_p.alignment = PP_ALIGN.CENTER
    ft_p.font.name = "Segoe UI"
    ft_p.font.size = Pt(10)
    ft_p.font.color.rgb = C_MUTED

def add_title_slide(prs, class_badge, main_title, sub_desc, topics_list):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_NAVY
    bg.line.fill.background()

    # Top Brand Ribbon
    ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.5))
    ribbon.fill.solid()
    ribbon.fill.fore_color.rgb = C_TEAL
    ribbon.line.fill.background()
    rtf = ribbon.text_frame
    rp = rtf.paragraphs[0]
    rp.text = "🎯 राष्ट्रीय आय-सह-मेधा छात्रवृत्ति परीक्षा (NMMS All India Special) • मिशन 35/35 अंक"
    rp.alignment = PP_ALIGN.CENTER
    rp.font.name = "Segoe UI"
    rp.font.size = Pt(12)
    rp.font.bold = True
    rp.font.color.rgb = C_WHITE

    # Logo
    if os.path.exists(LOGO_PATH):
        try:
            slide.shapes.add_picture(LOGO_PATH, Inches(6.0), Inches(0.7), Inches(1.33), Inches(1.33))
        except Exception:
            pass

    # Institute Name
    inst_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.05), Inches(12.333), Inches(0.6))
    itf = inst_box.text_frame
    ip = itf.paragraphs[0]
    ip.text = "SAGAR COACHING CENTRE, BHAGWANPUR (SUPAUL, BIHAR)"
    ip.alignment = PP_ALIGN.CENTER
    ip.font.name = "Segoe UI"
    ip.font.size = Pt(16)
    ip.font.bold = True
    ip.font.color.rgb = C_GOLD

    # Badge Card
    badge_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.66), Inches(2.7), Inches(6.0), Inches(0.5))
    badge_card.fill.solid()
    badge_card.fill.fore_color.rgb = C_TEAL
    badge_card.line.color.rgb = C_GOLD
    badge_card.line.width = Pt(1.5)
    btf = badge_card.text_frame
    bp = btf.paragraphs[0]
    bp.text = f"🌟 {class_badge} 🌟"
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Segoe UI"
    bp.font.size = Pt(13)
    bp.font.bold = True
    bp.font.color.rgb = C_WHITE

    # Main Title Box
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.25), Inches(12.333), Inches(1.2))
    ttf = title_box.text_frame
    ttf.word_wrap = True
    tp = ttf.paragraphs[0]
    tp.text = main_title
    tp.alignment = PP_ALIGN.CENTER
    tp.font.name = "Segoe UI"
    tp.font.size = Pt(30)
    tp.font.bold = True
    tp.font.color.rgb = C_WHITE

    tp_sub = ttf.add_paragraph()
    tp_sub.text = sub_desc
    tp_sub.alignment = PP_ALIGN.CENTER
    tp_sub.font.name = "Segoe UI"
    tp_sub.font.size = Pt(15)
    tp_sub.font.color.rgb = RGBColor(203, 213, 225)
    tp_sub.space_before = Pt(4)

    # Topics Card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.55), Inches(10.333), Inches(2.0))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(30, 41, 59)
    card.line.color.rgb = C_TEAL
    card.line.width = Pt(1.5)
    ctf = card.text_frame
    ctf.word_wrap = True
    ctf.margin_top = Inches(0.15)
    cp0 = ctf.paragraphs[0]
    cp0.text = "🎯 आज की क्लास में क्या सीखेंगे (Key Topics):"
    cp0.font.name = "Segoe UI"
    cp0.font.size = Pt(14)
    cp0.font.bold = True
    cp0.font.color.rgb = C_GOLD

    for t in topics_list:
        cp = ctf.add_paragraph()
        cp.text = f"  ✓  {t}"
        cp.font.name = "Segoe UI"
        cp.font.size = Pt(12)
        cp.font.color.rgb = C_WHITE
        cp.space_before = Pt(2)

    # Bottom Footer on Title
    bft = slide.shapes.add_textbox(Inches(0.5), Inches(6.75), Inches(12.333), Inches(0.5))
    bft_tf = bft.text_frame
    bp = bft_tf.paragraphs[0]
    bp.text = "मार्गदर्शक: सागर सर (NMMS King)  |  Helpline / WhatsApp: +91 91101 13671  |  Website: sagarcoaching.tech"
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Segoe UI"
    bp.font.size = Pt(11)
    bp.font.color.rgb = RGBColor(148, 163, 184)

def add_theory_slide(prs, title, points, slide_num=1, total_slides=1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, title, slide_num, total_slides)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = RGBColor(203, 213, 225)
    card.line.width = Pt(1.5)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.35)
    tf.margin_right = Inches(0.5)

    p0 = tf.paragraphs[0]
    p0.text = "📖 मुख्य वैज्ञानिक तथ्य एवं सिद्धांत (Key Concepts & Exam Notes):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(18)
    p0.font.bold = True
    p0.font.color.rgb = C_NAVY

    for pt in points:
        p = tf.add_paragraph()
        p.text = f"•  {pt}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(15)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(9)
        p.space_after = Pt(2)

def add_interactive_mcq_pair(prs, q_num, question, options, correct_opt, explanation, exam_tag, slide_num=1, total_slides=1, header_prefix="PYQ अभ्यास प्रश्न"):
    opt_labels = ["A", "B", "C", "D"]
    opt_pos = [
        (Inches(0.5), Inches(2.9), Inches(5.95), Inches(1.1)),
        (Inches(6.88), Inches(2.9), Inches(5.95), Inches(1.1)),
        (Inches(0.5), Inches(4.15), Inches(5.95), Inches(1.1)),
        (Inches(6.88), Inches(4.15), Inches(5.95), Inches(1.1)),
    ]

    # SLIDE 1: QUESTION ONLY (ANSWER HIDDEN)
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(s1, f"{header_prefix} #{q_num}  [{exam_tag}]", slide_num, total_slides)

    q_card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(1.5))
    q_card1.fill.solid()
    q_card1.fill.fore_color.rgb = C_BLUE_CARD
    q_card1.line.color.rgb = C_TEAL
    q_card1.line.width = Pt(2)

    q_tf1 = q_card1.text_frame
    q_tf1.word_wrap = True
    q_tf1.margin_left = Inches(0.3)
    q_tf1.vertical_anchor = MSO_ANCHOR.MIDDLE
    qp1 = q_tf1.paragraphs[0]
    qp1.text = f"प्र. {q_num}. {question}"
    qp1.font.name = "Segoe UI"
    qp1.font.size = Pt(18)
    qp1.font.bold = True
    qp1.font.color.rgb = C_NAVY

    for idx, (lbl, opt_text) in enumerate(zip(opt_labels, options)):
        left, top, w, h = opt_pos[idx]
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = RGBColor(203, 213, 225)
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = f"({lbl})  {opt_text}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_DARK

    # SLIDE 2: REVEAL SLIDE (ANSWER GLOWS GREEN + EXPLANATION)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(s2, f"{header_prefix} #{q_num}  [{exam_tag}]  —  [उत्तर व व्याख्या]", slide_num + 1, total_slides)

    q_card2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(1.5))
    q_card2.fill.solid()
    q_card2.fill.fore_color.rgb = C_BLUE_CARD
    q_card2.line.color.rgb = C_TEAL
    q_card2.line.width = Pt(2)

    q_tf2 = q_card2.text_frame
    q_tf2.word_wrap = True
    q_tf2.margin_left = Inches(0.3)
    q_tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
    qp2 = q_tf2.paragraphs[0]
    qp2.text = f"प्र. {q_num}. {question}"
    qp2.font.name = "Segoe UI"
    qp2.font.size = Pt(18)
    qp2.font.bold = True
    qp2.font.color.rgb = C_NAVY

    for idx, (lbl, opt_text) in enumerate(zip(opt_labels, options)):
        left, top, w, h = opt_pos[idx]
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        is_correct = (lbl == correct_opt)

        if is_correct:
            card.fill.fore_color.rgb = RGBColor(220, 252, 231)  # Soft emerald green
            card.line.color.rgb = C_GREEN
            card.line.width = Pt(3)
        else:
            card.fill.fore_color.rgb = C_WHITE
            card.line.color.rgb = RGBColor(226, 232, 240)
            card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        prefix = "✅ " if is_correct else ""
        p.text = f"{prefix}({lbl})  {opt_text}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_GREEN if is_correct else RGBColor(148, 163, 184)

    # Bottom Explanation Card
    ans_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(5.4), Inches(12.333), Inches(1.4))
    ans_card.fill.solid()
    ans_card.fill.fore_color.rgb = C_AMBER_CARD
    ans_card.line.color.rgb = C_GOLD
    ans_card.line.width = Pt(1.5)

    atf = ans_card.text_frame
    atf.word_wrap = True
    atf.margin_left = Inches(0.3)
    atf.margin_top = Inches(0.12)

    ap0 = atf.paragraphs[0]
    ap0.text = f"✅ सही उत्तर: विकल्प ({correct_opt})"
    ap0.font.name = "Segoe UI"
    ap0.font.size = Pt(15.5)
    ap0.font.bold = True
    ap0.font.color.rgb = C_GREEN

    ap1 = atf.add_paragraph()
    ap1.text = f"💡 व्याख्या (Explanation): {explanation}"
    ap1.font.name = "Segoe UI"
    ap1.font.size = Pt(14)
    ap1.font.bold = False
    ap1.font.color.rgb = C_DARK
    ap1.space_before = Pt(4)

def add_memory_key_slide(prs, title, key_items, slide_num=1, total_slides=1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, title, slide_num, total_slides)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card.fill.solid()
    card.fill.fore_color.rgb = C_AMBER_CARD
    card.line.color.rgb = C_GOLD
    card.line.width = Pt(2)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.35)

    p0 = tf.paragraphs[0]
    p0.text = "🔑 1-मिनट क्विक मेमोरी सूत्र (Revision Formulae):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(19)
    p0.font.bold = True
    p0.font.color.rgb = C_NAVY

    for item in key_items:
        p = tf.add_paragraph()
        p.text = f"⚡  {item}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = RGBColor(120, 53, 15)  # Dark amber
        p.space_before = Pt(12)
        p.space_after = Pt(2)

def add_homework_slide(prs, hw_questions, slide_num=1, total_slides=1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, "आज का अभ्यास कार्य (Homework Questions)", slide_num, total_slides)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card.fill.solid()
    card.fill.fore_color.rgb = C_BLUE_CARD
    card.line.color.rgb = C_TEAL
    card.line.width = Pt(2)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.35)

    p0 = tf.paragraphs[0]
    p0.text = "📝 घर पर अभ्यास के लिए 3 महत्वपूर्ण प्रश्न (Self Practice):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(19)
    p0.font.bold = True
    p0.font.color.rgb = C_NAVY

    for idx, hw in enumerate(hw_questions):
        p = tf.add_paragraph()
        p.text = f"प्रश्न {idx + 1}.  {hw}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(16)
        p.space_after = Pt(4)

def finalize_and_save_presentation(prs, filename, target_dir):
    actual_total = len(prs.slides)
    for idx, s in enumerate(prs.slides):
        for shape in s.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if "Slide " in p.text and " of " in p.text:
                        p.text = f"WhatsApp / Helpline: +91 91101 13671  |  वेबसाइट: sagarcoaching.tech  |  Slide {idx + 1} of {actual_total}"

    out_path = os.path.join(target_dir, filename)
    try:
        prs.save(out_path)
        print(f"--> [SAVED] {filename} ({actual_total} slides)")
    except PermissionError:
        temp_path = os.path.join(target_dir, filename.replace(".pptx", "_Updated.pptx"))
        prs.save(temp_path)
        print(f"--> [LOCKED - SAVED AS] {os.path.basename(temp_path)} ({actual_total} slides)")

if __name__ == "__main__":
    print("Science Base Module initialized successfully.")
