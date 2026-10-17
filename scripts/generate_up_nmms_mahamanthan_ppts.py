import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = r"d:\Work\E-Learning"
BASE_DIR = r"C:\Users\Orange\Downloads\New folder"
TARGET_DIR = os.path.join(BASE_DIR, "UP_NMMS_MahaManthan_Nov_Exam_SagarSir")
os.makedirs(TARGET_DIR, exist_ok=True)

LOGO_PATH = os.path.join(ROOT_DIR, "public", "logo-circle-transparent.png")
if not os.path.exists(LOGO_PATH):
    LOGO_PATH = os.path.join(BASE_DIR, "assets", "logo-circle-transparent.png")

# Import Questions
from data_up_mat_100q import UP_MAT_QUESTIONS
from data_up_sat_100q import UP_SAT_QUESTIONS

# Colors
C_NAVY = RGBColor(15, 23, 42)          # Deep Slate/Navy Primary
C_INDIGO = RGBColor(67, 56, 202)       # Royal Indigo
C_GOLD = RGBColor(217, 119, 6)         # Amber / Gold
C_LIGHT_BG = RGBColor(248, 250, 252)   # Soft Light Background
C_WHITE = RGBColor(255, 255, 255)
C_DARK = RGBColor(30, 41, 59)          # Dark Body Text
C_MUTED = RGBColor(100, 116, 139)      # Muted Text
C_BLUE_CARD = RGBColor(240, 249, 255)  # Light Sky Card
C_AMBER_CARD = RGBColor(254, 243, 199) # Soft Amber Card
C_GREEN = RGBColor(22, 101, 52)        # Forest Green Glow for Correct Answer (#166534)
C_RED = RGBColor(185, 28, 28)

def create_base_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_header_footer(slide, title_text, slide_num, total_slides):
    # Header Card
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.2), Inches(12.533), Inches(0.85)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = C_NAVY
    header.line.color.rgb = C_INDIGO
    header.line.width = Pt(1.5)

    # Logo in header (0.55 in to 1.30 in)
    if os.path.exists(LOGO_PATH):
        try:
            slide.shapes.add_picture(LOGO_PATH, Inches(0.55), Inches(0.25), Inches(0.75), Inches(0.75))
        except Exception:
            pass

    # Header Text in separate textbox
    tb = slide.shapes.add_textbox(Inches(1.45), Inches(0.2), Inches(11.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)

    p0 = tf.paragraphs[0]
    p0.text = "SAGAR COACHING CENTRE • UP NMMS नवम्बर परीक्षा 2024-25 महामंथन"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = C_GOLD

    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(13.5)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE

    # Footer
    footer_text = f"हेल्पलाइन / WhatsApp: +91 91101 13671  |  वेबसाइट: sagarcoaching.tech  |  Slide {slide_num} of {total_slides}"
    ft_box = slide.shapes.add_textbox(Inches(0.4), Inches(7.05), Inches(12.533), Inches(0.35))
    ft_tf = ft_box.text_frame
    ft_p = ft_tf.paragraphs[0]
    ft_p.text = footer_text
    ft_p.alignment = PP_ALIGN.CENTER
    ft_p.font.name = "Segoe UI"
    ft_p.font.size = Pt(10)
    ft_p.font.color.rgb = C_MUTED

def add_title_slide(prs, badge_text, main_title, sub_desc, topics_list):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_NAVY
    bg.line.fill.background()

    # Top Brand Ribbon
    ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.5))
    ribbon.fill.solid()
    ribbon.fill.fore_color.rgb = C_INDIGO
    ribbon.line.fill.background()
    rtf = ribbon.text_frame
    rp = rtf.paragraphs[0]
    rp.text = "🎯 राष्ट्रीय आय एवं योग्यता आधारित छात्रवृत्ति परीक्षा (UP NMMS 2024-25) • महामंथन मैराथन"
    rp.alignment = PP_ALIGN.CENTER
    rp.font.name = "Segoe UI"
    rp.font.size = Pt(12)
    rp.font.bold = True
    rp.font.color.rgb = C_WHITE

    # Logo
    if os.path.exists(LOGO_PATH):
        try:
            slide.shapes.add_picture(LOGO_PATH, Inches(6.0), Inches(0.65), Inches(1.35), Inches(1.35))
        except Exception:
            pass

    # Institute Name
    inst_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.02), Inches(12.333), Inches(0.6))
    itf = inst_box.text_frame
    ip = itf.paragraphs[0]
    ip.text = "SAGAR COACHING CENTRE, BHAGWANPUR"
    ip.alignment = PP_ALIGN.CENTER
    ip.font.name = "Segoe UI"
    ip.font.size = Pt(16)
    ip.font.bold = True
    ip.font.color.rgb = C_GOLD

    # Badge Card
    badge_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.2), Inches(2.65), Inches(6.933), Inches(0.5))
    badge_card.fill.solid()
    badge_card.fill.fore_color.rgb = C_INDIGO
    badge_card.line.color.rgb = C_GOLD
    badge_card.line.width = Pt(1.5)
    btf = badge_card.text_frame
    bp = btf.paragraphs[0]
    bp.text = f"🔥 {badge_text} 🔥"
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Segoe UI"
    bp.font.size = Pt(13)
    bp.font.bold = True
    bp.font.color.rgb = C_WHITE

    # Main Title Box
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(12.333), Inches(1.2))
    ttf = title_box.text_frame
    ttf.word_wrap = True
    tp = ttf.paragraphs[0]
    tp.text = main_title
    tp.alignment = PP_ALIGN.CENTER
    tp.font.name = "Segoe UI"
    tp.font.size = Pt(28)
    tp.font.bold = True
    tp.font.color.rgb = C_WHITE

    tp_sub = ttf.add_paragraph()
    tp_sub.text = sub_desc
    tp_sub.alignment = PP_ALIGN.CENTER
    tp_sub.font.name = "Segoe UI"
    tp_sub.font.size = Pt(14)
    tp_sub.font.color.rgb = RGBColor(203, 213, 225)
    tp_sub.space_before = Pt(4)

    # Topics Card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.5), Inches(10.333), Inches(2.1))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(30, 41, 59)
    card.line.color.rgb = C_INDIGO
    card.line.width = Pt(1.5)
    ctf = card.text_frame
    ctf.word_wrap = True
    ctf.margin_top = Inches(0.12)
    cp0 = ctf.paragraphs[0]
    cp0.text = "🎯 इस महामंथन मैराथन की मुख्य विशेषताएं (Live Class Highlights):"
    cp0.font.name = "Segoe UI"
    cp0.font.size = Pt(13.5)
    cp0.font.bold = True
    cp0.font.color.rgb = C_GOLD

    for t in topics_list:
        cp = ctf.add_paragraph()
        cp.text = f"  ✓  {t}"
        cp.font.name = "Segoe UI"
        cp.font.size = Pt(11.5)
        cp.font.color.rgb = C_WHITE
        cp.space_before = Pt(2)

    # Bottom Footer on Title
    bft = slide.shapes.add_textbox(Inches(0.5), Inches(6.75), Inches(12.333), Inches(0.5))
    bft_tf = bft.text_frame
    bp = bft_tf.paragraphs[0]
    bp.text = "मार्गदर्शक: सागर सर (NMMS Expert)  |  WhatsApp: +91 91101 13671  |  YouTube: Sagar Coaching Centre"
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Segoe UI"
    bp.font.size = Pt(11)
    bp.font.color.rgb = RGBColor(148, 163, 184)

def add_instructions_slide(prs, test_title, exam_type, total_slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, f"{test_title} — परीक्षा पैटर्न एवं लाइव क्लास निर्देश", 2, total_slides)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card.fill.solid()
    card.fill.fore_color.rgb = C_BLUE_CARD
    card.line.color.rgb = C_INDIGO
    card.line.width = Pt(2)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.3)

    p0 = tf.paragraphs[0]
    p0.text = f"📋 उत्तर प्रदेश NMMS {exam_type} परीक्षा योजना (Official Guidelines):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(18)
    p0.font.bold = True
    p0.font.color.rgb = C_NAVY

    if "MAT" in exam_type:
        rules = [
            "परीक्षा संरचना: मानसिक योग्यता परीक्षण (MAT) में कुल 90 बहुविकल्पीय प्रश्न (90 अंक) पूछे जाते हैं।",
            "इस महामंथन क्लास में UP NMMS के पिछले 10 वर्षों के ट्रेंड के आधार पर 102 अति-संभावित प्रश्न शामिल हैं।",
            "कवर किए गए प्रमुख विषय: श्रृंखला, सादृश्यता, वर्गीकरण, कोडिंग, लुप्त पद, गणितीय संक्रियाएं, रक्त संबंध, दिशा, वेन आरेख आदि।",
            "2-स्लाइड इंटरैक्टिव पद्धति: प्रश्न स्लाइड पर उत्तर पूरी तरह छुपा है ताकि लाइव क्लास में विद्यार्थी पहले स्वयं उत्तर सोचें व अभ्यास करें।",
            "अगली स्लाइड पर सही विकल्प हरे रंग (Green Glow) में चमकेगा और नीचे आसान ट्रिक व व्याख्या प्रस्तुत होगी।",
            "लक्ष्य: MAT में 90 में से 85+ अंक प्राप्त करके राज्य स्तरीय मेरिट में शीर्ष स्थान सुनिश्चित करना!"
        ]
    else:
        rules = [
            "परीक्षा संरचना: शैक्षिक योग्यता परीक्षण (SAT) में कुल 90 प्रश्न होते हैं (विज्ञान 35, सामाजिक विज्ञान 35, गणित 20)।",
            "इस महामंथन क्लास में UP SCERT कक्षा 7 व 8 बेसिक शिक्षा परिषद् के पूरे पाठ्यक्रम से 100 अति-महत्वपूर्ण प्रश्न शामिल हैं।",
            "विषय विभाजन: विज्ञान = 40 प्रश्न (भौतिक, रसायन, जीव) | सामाजिक विज्ञान = 40 प्रश्न (इतिहास, भूगोल, नागरिक) | गणित = 20 प्रश्न।",
            "उत्तर प्रदेश विशेष: दुधवा राष्ट्रीय उद्यान, नरोरा परमाणु संयंत्र, 1857 क्रांति के यूपी केंद्र, चौरी-चौरा, काकोरी आदि विशेष प्रश्न शामिल हैं।",
            "2-स्लाइड पद्धति: छात्र पहले प्रश्न देखकर विचार करेंगे, फिर अगली स्लाइड पर संपूर्ण व्याख्या व सूत्र समझेंगे।",
            "लक्ष्य: SAT में 90 में से 85+ अंक लाकर ₹48,000 की राष्ट्रीय छात्रवृत्ति पक्की करना!"
        ]

    for r in rules:
        p = tf.add_paragraph()
        p.text = f"•  {r}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(8)
        p.space_after = Pt(2)

def add_section_divider(prs, section_title, section_desc, slide_num, total_slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, f"विषय खंड — {section_title}", slide_num, total_slides)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.0), Inches(10.333), Inches(4.2))
    card.fill.solid()
    card.fill.fore_color.rgb = C_NAVY
    card.line.color.rgb = C_GOLD
    card.line.width = Pt(2)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.5)

    p0 = tf.paragraphs[0]
    p0.text = "⚡ आगामी महत्वपूर्ण खंड (Next Section) ⚡"
    p0.alignment = PP_ALIGN.CENTER
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(16)
    p0.font.bold = True
    p0.font.color.rgb = C_GOLD

    p1 = tf.add_paragraph()
    p1.text = section_title
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    p1.space_before = Pt(14)

    p2 = tf.add_paragraph()
    p2.text = section_desc
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(14.5)
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_before = Pt(10)

def add_interactive_mcq_pair(prs, q_num, q_hi, q_en, options, correct_opt, explanation, exam_tag, subject_name, slide_num, total_slides):
    opt_labels = ["A", "B", "C", "D"]
    opt_pos = [
        (Inches(0.5), Inches(2.95), Inches(5.95), Inches(1.1)),
        (Inches(6.88), Inches(2.95), Inches(5.95), Inches(1.1)),
        (Inches(0.5), Inches(4.18), Inches(5.95), Inches(1.1)),
        (Inches(6.88), Inches(4.18), Inches(5.95), Inches(1.1)),
    ]

    # SLIDE 1: QUESTION ONLY (ANSWER HIDDEN)
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(s1, f"{subject_name} • प्र. #{q_num}  [{exam_tag}]", slide_num, total_slides)

    q_card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.15), Inches(12.333), Inches(1.68))
    q_card1.fill.solid()
    q_card1.fill.fore_color.rgb = C_BLUE_CARD
    q_card1.line.color.rgb = C_INDIGO
    q_card1.line.width = Pt(2)

    q_tf1 = q_card1.text_frame
    q_tf1.word_wrap = True
    q_tf1.margin_left = Inches(0.3)
    q_tf1.margin_top = Inches(0.1)
    q_tf1.margin_bottom = Inches(0.1)
    q_tf1.vertical_anchor = MSO_ANCHOR.MIDDLE

    qp1_hi = q_tf1.paragraphs[0]
    qp1_hi.text = f"प्र. {q_num}. {q_hi}"
    qp1_hi.alignment = PP_ALIGN.CENTER
    qp1_hi.font.name = "Segoe UI"
    qp1_hi.font.size = Pt(16.5)
    qp1_hi.font.bold = True
    qp1_hi.font.color.rgb = C_NAVY

    qp1_en = q_tf1.add_paragraph()
    qp1_en.text = f"Q. {q_num}. {q_en}"
    qp1_en.alignment = PP_ALIGN.CENTER
    qp1_en.font.name = "Segoe UI"
    qp1_en.font.size = Pt(14.5)
    qp1_en.font.bold = True
    qp1_en.font.color.rgb = RGBColor(37, 99, 235)
    qp1_en.space_before = Pt(6)

    for idx, (lbl, opt_text) in enumerate(zip(opt_labels, options)):
        left, top, w, h = opt_pos[idx]
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = RGBColor(203, 213, 225)
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = f"({lbl})  {opt_text}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_DARK

    # SLIDE 2: REVEAL SLIDE (ANSWER GLOWS GREEN + EXPLANATION)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(s2, f"{subject_name} • प्र. #{q_num}  [{exam_tag}]  —  [उत्तर व व्याख्या]", slide_num + 1, total_slides)

    q_card2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.15), Inches(12.333), Inches(1.68))
    q_card2.fill.solid()
    q_card2.fill.fore_color.rgb = C_BLUE_CARD
    q_card2.line.color.rgb = C_INDIGO
    q_card2.line.width = Pt(2)

    q_tf2 = q_card2.text_frame
    q_tf2.word_wrap = True
    q_tf2.margin_left = Inches(0.3)
    q_tf2.margin_top = Inches(0.1)
    q_tf2.margin_bottom = Inches(0.1)
    q_tf2.vertical_anchor = MSO_ANCHOR.MIDDLE

    qp2_hi = q_tf2.paragraphs[0]
    qp2_hi.text = f"प्र. {q_num}. {q_hi}"
    qp2_hi.alignment = PP_ALIGN.CENTER
    qp2_hi.font.name = "Segoe UI"
    qp2_hi.font.size = Pt(16.5)
    qp2_hi.font.bold = True
    qp2_hi.font.color.rgb = C_NAVY

    qp2_en = q_tf2.add_paragraph()
    qp2_en.text = f"Q. {q_num}. {q_en}"
    qp2_en.alignment = PP_ALIGN.CENTER
    qp2_en.font.name = "Segoe UI"
    qp2_en.font.size = Pt(14.5)
    qp2_en.font.bold = True
    qp2_en.font.color.rgb = RGBColor(37, 99, 235)
    qp2_en.space_before = Pt(6)

    for idx, (lbl, opt_text) in enumerate(zip(opt_labels, options)):
        left, top, w, h = opt_pos[idx]
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        is_correct = (lbl == correct_opt)

        if is_correct:
            card.fill.fore_color.rgb = RGBColor(220, 252, 231)  # Emerald light
            card.line.color.rgb = C_GREEN
            card.line.width = Pt(3)
        else:
            card.fill.fore_color.rgb = C_WHITE
            card.line.color.rgb = RGBColor(226, 232, 240)
            card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        prefix = "✅ " if is_correct else ""
        p.text = f"{prefix}({lbl})  {opt_text}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_GREEN if is_correct else RGBColor(148, 163, 184)

    # Bottom Explanation Card
    ans_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(5.42), Inches(12.333), Inches(1.5))
    ans_card.fill.solid()
    ans_card.fill.fore_color.rgb = C_AMBER_CARD
    ans_card.line.color.rgb = C_GOLD
    ans_card.line.width = Pt(1.5)

    atf = ans_card.text_frame
    atf.word_wrap = True
    atf.margin_left = Inches(0.3)
    atf.margin_top = Inches(0.12)

    ap0 = atf.paragraphs[0]
    ap0.text = f"✅ सही उत्तर / Correct Option: ({correct_opt})"
    ap0.font.name = "Segoe UI"
    ap0.font.size = Pt(15)
    ap0.font.bold = True
    ap0.font.color.rgb = C_GREEN

    ap1 = atf.add_paragraph()
    ap1.text = f"💡 व्याख्या एवं ट्रिक (Explanation): {explanation}"
    ap1.font.name = "Segoe UI"
    ap1.font.size = Pt(13)
    ap1.font.bold = False
    ap1.font.color.rgb = C_DARK
    ap1.space_before = Pt(3)

def add_outro_slide(prs, exam_title, total_slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, f"{exam_title} — सफलता का संकल्प एवं संपर्क सूत्र", total_slides, total_slides)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.333), Inches(5.2))
    card.fill.solid()
    card.fill.fore_color.rgb = C_NAVY
    card.line.color.rgb = C_GOLD
    card.line.width = Pt(2.5)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.4)

    p0 = tf.paragraphs[0]
    p0.text = "🏆 आपकी सफलता ही हमारा संकल्प है! 🏆"
    p0.alignment = PP_ALIGN.CENTER
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(22)
    p0.font.bold = True
    p0.font.color.rgb = C_GOLD

    p1 = tf.add_paragraph()
    p1.text = "SAGAR COACHING CENTRE, BHAGWANPUR"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    p1.space_before = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = "उत्तर प्रदेश NMMS छात्रवृत्ति परीक्षा (₹48,000) में चयनित होने वाले सभी भावी मेधावियों को अग्रिम शुभकामनाएं!"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(226, 232, 240)
    p2.space_before = Pt(10)

    p3 = tf.add_paragraph()
    p3.text = "📞 ऑनलाइन टेस्ट, PDF नोट्स एवं सहायता हेतु: +91 91101 13671\n🌐 वेबसाइट: sagarcoaching.tech  |  📺 YouTube: Sagar Coaching Centre"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(15)
    p3.font.bold = True
    p3.font.color.rgb = C_GOLD
    p3.space_before = Pt(16)


def build_up_mat_deck():
    print("Generating UP NMMS MAT MahaManthan Deck (Bilingual)...")
    prs = create_base_presentation()

    total_questions = len(UP_MAT_QUESTIONS)
    # Total slides = 1 (Title) + 1 (Instructions) + (102 * 2 = 204) + 1 (Outro) = 207 slides
    total_slides = 2 + (total_questions * 2) + 1

    add_title_slide(
        prs,
        badge_text="UP NMMS नवम्बर परीक्षा 2024-25 महामंथन (Bilingual)",
        main_title="मानसिक योग्यता परीक्षण (MAT) 100+ सुपर मैराथन",
        sub_desc="Mental Ability Test (MAT) • 100+ Most Predicted Questions (Hindi & English)",
        topics_list=[
            "संख्या व अक्षर श्रेणी (Number & Letter Series)",
            "सादृश्यता एवं वर्गीकरण (Analogy & Odd-One Out)",
            "कोडिंग-डिकोडिंग एवं लुप्त संख्याएं (Missing Numbers in Figures)",
            "गणितीय संक्रियाएं, रक्त संबंध, दिशा-दूरी एवं वेन आरेख (Bilingual Format)"
        ]
    )

    add_instructions_slide(prs, "UP NMMS MAT महामंथन", "MAT (मानसिक योग्यता)", total_slides)

    curr_slide = 3
    for q in UP_MAT_QUESTIONS:
        add_interactive_mcq_pair(
            prs=prs,
            q_num=q["q_num"],
            q_hi=q["question_hi"],
            q_en=q["question_en"],
            options=q["options"],
            correct_opt=q["correct"],
            explanation=q["explanation"],
            exam_tag="UP NMMS Bilingual",
            subject_name=f"MAT • {q['topic']}",
            slide_num=curr_slide,
            total_slides=total_slides
        )
        curr_slide += 2

    add_outro_slide(prs, "UP NMMS MAT महामंथन", total_slides)

    out_file = os.path.join(TARGET_DIR, "UP_NMMS_MAT_MahaManthan_100Q_Nov_Exam.pptx")
    prs.save(out_file)
    print(f"MAT Deck Saved: {out_file} (Total Slides: {len(prs.slides)})")
    return out_file


def build_up_sat_deck():
    print("Generating UP NMMS SAT MahaManthan Deck (Bilingual)...")
    prs = create_base_presentation()

    total_questions = len(UP_SAT_QUESTIONS)
    # Total slides = 1 (Title) + 1 (Instructions) + (100 * 2 = 200) + 1 (Outro) = 203 slides
    total_slides = 2 + (total_questions * 2) + 1

    add_title_slide(
        prs,
        badge_text="UP NMMS नवम्बर परीक्षा 2024-25 महामंथन (Bilingual)",
        main_title="शैक्षिक योग्यता परीक्षण (SAT) 100 प्रश्न महा मैराथन",
        sub_desc="Science (40Q) + Social Studies (40Q) + Mathematics (20Q) • Hindi & English",
        topics_list=[
            "विज्ञान / Science: भौतिकी (14Q), रसायन (13Q), जीव विज्ञान (13Q)",
            "सामाजिक अध्ययन / SST: इतिहास (15Q), भूगोल (15Q), नागरिक शास्त्र (10Q)",
            "गणित / Mathematics: अंकगणित, बीजगणित, रेखागणित, क्षेत्रमिति एवं सांख्यिकी (20Q)",
            "उत्तर प्रदेश विशेष: दुधवा पार्क, नरोरा संयंत्र, 1857 केंद्र, प्रयागराज संगम (द्विभाषी)"
        ]
    )

    add_instructions_slide(prs, "UP NMMS SAT महामंथन", "SAT (शैक्षिक योग्यता)", total_slides)

    curr_slide = 3
    for q in UP_SAT_QUESTIONS:
        add_interactive_mcq_pair(
            prs=prs,
            q_num=q["q_num"],
            q_hi=q["question_hi"],
            q_en=q["question_en"],
            options=q["options"],
            correct_opt=q["correct"],
            explanation=q["explanation"],
            exam_tag="UP NMMS Bilingual",
            subject_name=f"SAT • {q['subject']}",
            slide_num=curr_slide,
            total_slides=total_slides
        )
        curr_slide += 2

    add_outro_slide(prs, "UP NMMS SAT महामंथन", total_slides)

    out_file = os.path.join(TARGET_DIR, "UP_NMMS_SAT_MahaManthan_100Q_Nov_Exam.pptx")
    prs.save(out_file)
    print(f"SAT Deck Saved: {out_file} (Total Slides: {len(prs.slides)})")
    return out_file

if __name__ == "__main__":
    mat_ppt = build_up_mat_deck()
    sat_ppt = build_up_sat_deck()
    print("Both UP NMMS MahaManthan PPTs generated successfully!")
