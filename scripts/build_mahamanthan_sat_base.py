import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = r"d:\Work\E-Learning"
BASE_DIR = r"C:\Users\Orange\Downloads\New folder"
MAHAMANTHAN_SAT_DIR = os.path.join(BASE_DIR, "NMMS_SAT_Grand_MahaManthan_SagarSir")
os.makedirs(MAHAMANTHAN_SAT_DIR, exist_ok=True)

LOGO_PATH = os.path.join(ROOT_DIR, "public", "logo-circle-transparent.png")
if not os.path.exists(LOGO_PATH):
    LOGO_PATH = os.path.join(BASE_DIR, "assets", "logo-circle-transparent.png")

# Color Palette: Deep Royal Indigo / Navy + Gold Accent + Emerald Green Glow
C_NAVY = RGBColor(15, 23, 42)          # Deep Slate/Navy Primary
C_INDIGO = RGBColor(67, 56, 202)       # Royal Indigo
C_GOLD = RGBColor(217, 119, 6)         # Amber / Gold
C_LIGHT_BG = RGBColor(248, 250, 252)   # Soft Light Background
C_WHITE = RGBColor(255, 255, 255)
C_DARK = RGBColor(30, 41, 59)          # Dark Body Text
C_MUTED = RGBColor(100, 116, 139)      # Muted Text
C_BLUE_CARD = RGBColor(240, 249, 255)  # Light Sky Card
C_AMBER_CARD = RGBColor(254, 243, 199) # Soft Amber Card
C_GREEN = RGBColor(22, 101, 52)        # Forest Green Glow for Correct Answer (#166534)
C_RED = RGBColor(185, 28, 28)

def create_base_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_header_footer(slide, title_text, slide_num, total_slides):
    # Header Card
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.2), Inches(12.533), Inches(0.85)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = C_NAVY
    header.line.color.rgb = C_INDIGO
    header.line.width = Pt(1.5)

    # Logo in header (0.55 in to 1.30 in)
    if os.path.exists(LOGO_PATH):
        try:
            slide.shapes.add_picture(LOGO_PATH, Inches(0.55), Inches(0.25), Inches(0.75), Inches(0.75))
        except Exception:
            pass

    # Header Text in separate textbox (starts at 1.45 in, no text overlap)
    tb = slide.shapes.add_textbox(Inches(1.45), Inches(0.2), Inches(11.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)

    p0 = tf.paragraphs[0]
    p0.text = "SAGAR COACHING CENTRE BHAGWANPUR • NMMS All India (SAT महामंथन)"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = C_GOLD

    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(13.5)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE

    # Footer
    footer_text = f"WhatsApp / Helpline: +91 91101 13671  |  वेबसाइट: sagarcoaching.tech  |  Slide {slide_num} of {total_slides}"
    ft_box = slide.shapes.add_textbox(Inches(0.4), Inches(7.05), Inches(12.533), Inches(0.35))
    ft_tf = ft_box.text_frame
    ft_p = ft_tf.paragraphs[0]
    ft_p.text = footer_text
    ft_p.alignment = PP_ALIGN.CENTER
    ft_p.font.name = "Segoe UI"
    ft_p.font.size = Pt(10)
    ft_p.font.color.rgb = C_MUTED

def add_title_slide(prs, badge_text, main_title, sub_desc, topics_list):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_NAVY
    bg.line.fill.background()

    # Top Brand Ribbon
    ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.5))
    ribbon.fill.solid()
    ribbon.fill.fore_color.rgb = C_INDIGO
    ribbon.line.fill.background()
    rtf = ribbon.text_frame
    rp = rtf.paragraphs[0]
    rp.text = "🎯 राष्ट्रीय आय-सह-मेधा छात्रवृत्ति परीक्षा (NMMS All India) • सम्पूर्ण SAT 90/90 अंक महामंथन"
    rp.alignment = PP_ALIGN.CENTER
    rp.font.name = "Segoe UI"
    rp.font.size = Pt(12)
    rp.font.bold = True
    rp.font.color.rgb = C_WHITE

    # Logo
    if os.path.exists(LOGO_PATH):
        try:
            slide.shapes.add_picture(LOGO_PATH, Inches(6.0), Inches(0.7), Inches(1.33), Inches(1.33))
        except Exception:
            pass

    # Institute Name
    inst_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.05), Inches(12.333), Inches(0.6))
    itf = inst_box.text_frame
    ip = itf.paragraphs[0]
    ip.text = "SAGAR COACHING CENTRE, BHAGWANPUR (SUPAUL, BIHAR)"
    ip.alignment = PP_ALIGN.CENTER
    ip.font.name = "Segoe UI"
    ip.font.size = Pt(16)
    ip.font.bold = True
    ip.font.color.rgb = C_GOLD

    # Badge Card
    badge_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.4), Inches(2.7), Inches(6.533), Inches(0.5))
    badge_card.fill.solid()
    badge_card.fill.fore_color.rgb = C_INDIGO
    badge_card.line.color.rgb = C_GOLD
    badge_card.line.width = Pt(1.5)
    btf = badge_card.text_frame
    bp = btf.paragraphs[0]
    bp.text = f"🌟 {badge_text} 🌟"
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Segoe UI"
    bp.font.size = Pt(13)
    bp.font.bold = True
    bp.font.color.rgb = C_WHITE

    # Main Title Box
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.25), Inches(12.333), Inches(1.2))
    ttf = title_box.text_frame
    ttf.word_wrap = True
    tp = ttf.paragraphs[0]
    tp.text = main_title
    tp.alignment = PP_ALIGN.CENTER
    tp.font.name = "Segoe UI"
    tp.font.size = Pt(28)
    tp.font.bold = True
    tp.font.color.rgb = C_WHITE

    tp_sub = ttf.add_paragraph()
    tp_sub.text = sub_desc
    tp_sub.alignment = PP_ALIGN.CENTER
    tp_sub.font.name = "Segoe UI"
    tp_sub.font.size = Pt(14)
    tp_sub.font.color.rgb = RGBColor(203, 213, 225)
    tp_sub.space_before = Pt(4)

    # Topics Card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.55), Inches(10.333), Inches(2.0))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(30, 41, 59)
    card.line.color.rgb = C_INDIGO
    card.line.width = Pt(1.5)
    ctf = card.text_frame
    ctf.word_wrap = True
    ctf.margin_top = Inches(0.12)
    cp0 = ctf.paragraphs[0]
    cp0.text = "🎯 इस 90 प्रश्नों के महामंथन टेस्ट का पाठ्यक्रम विवरण (Syllabus Structure):"
    cp0.font.name = "Segoe UI"
    cp0.font.size = Pt(13.5)
    cp0.font.bold = True
    cp0.font.color.rgb = C_GOLD

    for t in topics_list:
        cp = ctf.add_paragraph()
        cp.text = f"  ✓  {t}"
        cp.font.name = "Segoe UI"
        cp.font.size = Pt(11.5)
        cp.font.color.rgb = C_WHITE
        cp.space_before = Pt(2)

    # Bottom Footer on Title
    bft = slide.shapes.add_textbox(Inches(0.5), Inches(6.75), Inches(12.333), Inches(0.5))
    bft_tf = bft.text_frame
    bp = bft_tf.paragraphs[0]
    bp.text = "मार्गदर्शक: सागर सर (NMMS King)  |  Helpline / WhatsApp: +91 91101 13671  |  Website: sagarcoaching.tech"
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Segoe UI"
    bp.font.size = Pt(11)
    bp.font.color.rgb = RGBColor(148, 163, 184)

def add_instructions_slide(prs, test_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, f"{test_title} — परीक्षा निर्देश एवं अंक योजना", 2, 185)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card.fill.solid()
    card.fill.fore_color.rgb = C_BLUE_CARD
    card.line.color.rgb = C_INDIGO
    card.line.width = Pt(2)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.3)

    p0 = tf.paragraphs[0]
    p0.text = "📋 NMMS SAT सम्पूर्ण 90 प्रश्न परीक्षा योजना (Official Exam Scheme):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(18)
    p0.font.bold = True
    p0.font.color.rgb = C_NAVY

    rules = [
        "विषय-वार अंक विभाजन: विज्ञान = 35 प्रश्न (35 अंक) | सामाजिक विज्ञान = 35 प्रश्न (35 अंक) | गणित = 20 प्रश्न (20 अंक)। कुल = 90 प्रश्न (90 अंक)।",
        "समय सीमा: वास्तविक परीक्षा में 90 प्रश्नों हेतु 90 मिनट का समय मिलता है। कोई नकारात्मक अंकन (No Negative Marking) नहीं है।",
        "2-स्लाइड इंटरैक्टिव पद्धति: प्रश्न स्लाइड पर उत्तर पूरी तरह छुपा हुआ है ताकि छात्र पहले स्वयं विचार कर जवाब दें।",
        "अगली स्लाइड पर सही विकल्प हरे रंग (Green Glow) में प्रकाशित होगा तथा नीचे संपूर्ण हल और व्याख्या दी गई है।",
        "लक्ष्य: NMMS परीक्षा में राज्य मेरिट सूची में शीर्ष स्थान प्राप्त करना और ₹48,000 की छात्रवृत्ति सुनिश्चित करना!"
    ]

    for r in rules:
        p = tf.add_paragraph()
        p.text = f"•  {r}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(14.5)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(10)
        p.space_after = Pt(2)

def add_interactive_mcq_pair(prs, q_num, question, options, correct_opt, explanation, exam_tag, subject_name, slide_num=1, total_slides=185):
    opt_labels = ["A", "B", "C", "D"]
    opt_pos = [
        (Inches(0.5), Inches(2.9), Inches(5.95), Inches(1.1)),
        (Inches(6.88), Inches(2.9), Inches(5.95), Inches(1.1)),
        (Inches(0.5), Inches(4.15), Inches(5.95), Inches(1.1)),
        (Inches(6.88), Inches(4.15), Inches(5.95), Inches(1.1)),
    ]

    # SLIDE 1: QUESTION ONLY (ANSWER HIDDEN)
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(s1, f"{subject_name} • प्र. #{q_num}  [{exam_tag}]", slide_num, total_slides)

    q_card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(1.5))
    q_card1.fill.solid()
    q_card1.fill.fore_color.rgb = C_BLUE_CARD
    q_card1.line.color.rgb = C_INDIGO
    q_card1.line.width = Pt(2)

    q_tf1 = q_card1.text_frame
    q_tf1.word_wrap = True
    q_tf1.margin_left = Inches(0.3)
    q_tf1.vertical_anchor = MSO_ANCHOR.MIDDLE
    qp1 = q_tf1.paragraphs[0]
    qp1.text = f"प्र. {q_num}. {question}"
    qp1.font.name = "Segoe UI"
    qp1.font.size = Pt(17.5)
    qp1.font.bold = True
    qp1.font.color.rgb = C_NAVY

    for idx, (lbl, opt_text) in enumerate(zip(opt_labels, options)):
        left, top, w, h = opt_pos[idx]
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = RGBColor(203, 213, 225)
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = f"({lbl})  {opt_text}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(15.5)
        p.font.bold = True
        p.font.color.rgb = C_DARK

    # SLIDE 2: REVEAL SLIDE (ANSWER GLOWS GREEN + EXPLANATION)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(s2, f"{subject_name} • प्र. #{q_num}  [{exam_tag}]  —  [उत्तर व व्याख्या]", slide_num + 1, total_slides)

    q_card2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(1.5))
    q_card2.fill.solid()
    q_card2.fill.fore_color.rgb = C_BLUE_CARD
    q_card2.line.color.rgb = C_INDIGO
    q_card2.line.width = Pt(2)

    q_tf2 = q_card2.text_frame
    q_tf2.word_wrap = True
    q_tf2.margin_left = Inches(0.3)
    q_tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
    qp2 = q_tf2.paragraphs[0]
    qp2.text = f"प्र. {q_num}. {question}"
    qp2.font.name = "Segoe UI"
    qp2.font.size = Pt(17.5)
    qp2.font.bold = True
    qp2.font.color.rgb = C_NAVY

    for idx, (lbl, opt_text) in enumerate(zip(opt_labels, options)):
        left, top, w, h = opt_pos[idx]
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        is_correct = (lbl == correct_opt)

        if is_correct:
            card.fill.fore_color.rgb = RGBColor(220, 252, 231)  # Emerald light
            card.line.color.rgb = C_GREEN
            card.line.width = Pt(3)
        else:
            card.fill.fore_color.rgb = C_WHITE
            card.line.color.rgb = RGBColor(226, 232, 240)
            card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        prefix = "✅ " if is_correct else ""
        p.text = f"{prefix}({lbl})  {opt_text}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(15.5)
        p.font.bold = True
        p.font.color.rgb = C_GREEN if is_correct else RGBColor(148, 163, 184)

    # Bottom Explanation Card
    ans_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(5.4), Inches(12.333), Inches(1.4))
    ans_card.fill.solid()
    ans_card.fill.fore_color.rgb = C_AMBER_CARD
    ans_card.line.color.rgb = C_GOLD
    ans_card.line.width = Pt(1.5)

    atf = ans_card.text_frame
    atf.word_wrap = True
    atf.margin_left = Inches(0.3)
    atf.margin_top = Inches(0.12)

    ap0 = atf.paragraphs[0]
    ap0.text = f"✅ सही उत्तर: विकल्प ({correct_opt})"
    ap0.font.name = "Segoe UI"
    ap0.font.size = Pt(15.5)
    ap0.font.bold = True
    ap0.font.color.rgb = C_GREEN

    ap1 = atf.add_paragraph()
    ap1.text = f"💡 व्याख्या (Explanation): {explanation}"
    ap1.font.name = "Segoe UI"
    ap1.font.size = Pt(14)
    ap1.font.bold = False
    ap1.font.color.rgb = C_DARK
    ap1.space_before = Pt(4)

def add_score_cutoff_slide(prs, test_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, f"{test_title} — स्कोर विश्लेषण व कट-ऑफ मैट्रिक्स", 183, 185)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = C_GOLD
    card.line.width = Pt(2)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.35)

    p0 = tf.paragraphs[0]
    p0.text = "🎯 90 प्रश्नों का स्कोर विश्लेषण एवं राज्य मेरिट रणनीति (Cut-off Strategy):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(18)
    p0.font.bold = True
    p0.font.color.rgb = C_NAVY

    rows = [
        "80+ अंक (State Top Rankers): राज्य में टॉप 50 रैंक सुनिश्चित! छात्रवृत्ति 100% पक्की।",
        "70 - 79 अंक (Very High Merit): सुरक्षित कट-ऑफ पार! जिला स्तर पर उत्कृष्ट रैंक।",
        "60 - 69 अंक (Qualifying Zone): गणित एवं रसायन विज्ञान के सूत्रों का 1 बार पुनः अभ्यास करें।",
        "60 से कम अंक: विज्ञान व सामाजिक विज्ञान के सभी 17+17 क्लास नोट्स का दोबारा रिवीजन करें।"
    ]

    for r in rows:
        p = tf.add_paragraph()
        p.text = f"🏆  {r}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(15)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(12)

def add_memory_key_slide(prs, test_title, keys_list):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, f"{test_title} — त्वरित ऑल-सब्जेक्ट गोल्डन सूत्र", 184, 185)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card.fill.solid()
    card.fill.fore_color.rgb = C_AMBER_CARD
    card.line.color.rgb = C_GOLD
    card.line.width = Pt(2)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.3)

    p0 = tf.paragraphs[0]
    p0.text = "🔑 1-मिनट क्विक ऑल-सब्जेक्ट गोल्डन रिवीजन सूत्र (Grand Memory Key):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(18)
    p0.font.bold = True
    p0.font.color.rgb = C_NAVY

    for k in keys_list:
        p = tf.add_paragraph()
        p.text = f"⚡  {k}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(14.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(120, 53, 15)
        p.space_before = Pt(10)
        p.space_after = Pt(2)

def add_closing_branding_slide(prs, badge_text, test_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, f"{test_title} — सफलता का संकल्प", 185, 185)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card.fill.solid()
    card.fill.fore_color.rgb = C_NAVY
    card.line.color.rgb = C_GOLD
    card.line.width = Pt(2.5)

    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p0 = tf.paragraphs[0]
    p0.text = "SAGAR COACHING CENTRE, BHAGWANPUR"
    p0.alignment = PP_ALIGN.CENTER
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(24)
    p0.font.bold = True
    p0.font.color.rgb = C_GOLD

    p1 = tf.add_paragraph()
    p1.text = "निकट - माँ भगवती मंदिर, भगवानपुर (सुपौल, बिहार) • पिन - 852109"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(15)
    p1.font.color.rgb = RGBColor(203, 213, 225)
    p1.space_before = Pt(6)

    p2 = tf.add_paragraph()
    p2.text = f"🌟 {badge_text} : 90/90 अंक सम्पूर्ण SAT महामंथन 🌟"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = C_WHITE
    p2.space_before = Pt(16)

    p3 = tf.add_paragraph()
    p3.text = "मार्गदर्शक: सागर सर (NMMS King)  |  WhatsApp / Helpline: +91 91101 13671"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(16)
    p3.font.bold = True
    p3.font.color.rgb = C_GOLD
    p3.space_before = Pt(14)

    p4 = tf.add_paragraph()
    p4.text = "ऑनलाइन टेस्ट, क्लास नोट्स व PDF सामग्री हेतु वेबसाइट पर पधारें: sagarcoaching.tech"
    p4.alignment = PP_ALIGN.CENTER
    p4.font.name = "Segoe UI"
    p4.font.size = Pt(14)
    p4.font.color.rgb = RGBColor(226, 232, 240)
    p4.space_before = Pt(8)

def finalize_and_save_sat_presentation(prs, filename):
    actual_total = len(prs.slides)
    for idx, s in enumerate(prs.slides):
        for shape in s.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if "Slide " in p.text and " of " in p.text:
                        p.text = f"WhatsApp / Helpline: +91 91101 13671  |  वेबसाइट: sagarcoaching.tech  |  Slide {idx + 1} of {actual_total}"

    out_path = os.path.join(MAHAMANTHAN_SAT_DIR, filename)
    prs.save(out_path)
    print(f"--> [SAVED] {filename} ({actual_total} slides)")

print("MahaManthan SAT Base ready.")
