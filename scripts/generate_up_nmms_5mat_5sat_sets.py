# -*- coding: utf-8 -*-
"""
Master Generator for UP NMMS 5 MAT and 5 SAT MahaManthan Practice Sets.
Creates:
- 5 Dedicated MAT Sets (Set 01 - Set 05) under UP_NMMS_MAT_MahaManthan_5_Sets_SagarSir/
- 5 Dedicated SAT Sets (Set 01 - Set 05) under UP_NMMS_SAT_MahaManthan_5_Sets_SagarSir/

Features:
- Exactly 50 questions per set (103 slides per presentation)
- 2-slide interactive format (Answer hidden on Slide 1; Emerald Green Glow + Amber Explanation on Slide 2)
- Questions centered with comfortable vertical spacing
- Bilingual (Hindi on line 1, English on line 2)
- Zero forbidden words (NO comment, timer, 10s)
"""

import os
import sys

# Force UTF-8 stdout
if sys.platform == "win32":
    import codecs
    sys.stdout = codecs.getwriter("utf-8")(sys.stdout.detach())

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Question Pools
import data_up_mat_100q as m1
import data_up_mat_additional as m2
import data_up_pool3 as p3
import data_up_pool4 as p4
import data_up_sat_100q as s1
import data_up_sat_additional as s2

# Base Paths
BASE_DIR = r"C:\Users\Orange\Downloads\New folder"
MAT_DIR = os.path.join(BASE_DIR, "UP_NMMS_MAT_MahaManthan_5_Sets_SagarSir")
SAT_DIR = os.path.join(BASE_DIR, "UP_NMMS_SAT_MahaManthan_5_Sets_SagarSir")
LOGO_PATH = os.path.join(BASE_DIR, "assets", "sagar_classes_logo.png")

# Premium Color Palette
C_NAVY = RGBColor(15, 23, 42)          # #0F172A Deep Navy
C_INDIGO = RGBColor(30, 58, 138)       # #1E3A8A Dark Indigo
C_GOLD = RGBColor(217, 119, 6)         # #D97706 Warm Amber Gold
C_WHITE = RGBColor(255, 255, 255)
C_DARK = RGBColor(15, 23, 42)
C_MUTED = RGBColor(100, 116, 139)
C_BLUE_CARD = RGBColor(241, 245, 249)  # Soft background
C_AMBER_CARD = RGBColor(254, 243, 199) # Explanation amber tint
C_GREEN = RGBColor(22, 101, 52)        # Emerald green #166534

TOTAL_QUESTIONS_PER_SET = 50
TOTAL_SLIDES_PER_DECK = 103  # Title (1) + Instructions (1) + 50*2 (100) + Outro (1) = 103

def create_base_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_header_footer(slide, title_text, slide_num, total_slides):
    # Top Ribbon
    ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.95))
    ribbon.fill.solid()
    ribbon.fill.fore_color.rgb = C_NAVY
    ribbon.line.fill.background()

    tf = ribbon.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.08)

    p0 = tf.paragraphs[0]
    p0.text = "SAGAR COACHING CENTRE • UP NMMS परीक्षा महामंथन अभ्यास सेट"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = C_GOLD

    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE

    # Logo
    if os.path.exists(LOGO_PATH):
        try:
            slide.shapes.add_picture(LOGO_PATH, Inches(12.35), Inches(0.08), Inches(0.8), Inches(0.8))
        except Exception:
            pass

    # Footer
    footer_text = f"हेल्पलाइन / WhatsApp: +91 91101 13671  |  वेबसाइट: sagarcoaching.tech  |  Slide {slide_num} of {total_slides}"
    ft_box = slide.shapes.add_textbox(Inches(0.4), Inches(7.05), Inches(12.533), Inches(0.35))
    ft_tf = ft_box.text_frame
    ft_p = ft_tf.paragraphs[0]
    ft_p.text = footer_text
    ft_p.alignment = PP_ALIGN.CENTER
    ft_p.font.name = "Segoe UI"
    ft_p.font.size = Pt(10)
    ft_p.font.color.rgb = C_MUTED

def add_title_slide(prs, badge_text, main_title, sub_desc, topics_list, set_label):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_NAVY
    bg.line.fill.background()

    # Top Brand Ribbon
    ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.5))
    ribbon.fill.solid()
    ribbon.fill.fore_color.rgb = C_INDIGO
    ribbon.line.fill.background()
    rtf = ribbon.text_frame
    rp = rtf.paragraphs[0]
    rp.text = "🎯 राष्ट्रीय आय एवं योग्यता आधारित छात्रवृत्ति परीक्षा (UP NMMS 2024-25) • महामंथन अभ्यास सेट"
    rp.alignment = PP_ALIGN.CENTER
    rp.font.name = "Segoe UI"
    rp.font.size = Pt(12)
    rp.font.bold = True
    rp.font.color.rgb = C_WHITE

    # Logo
    if os.path.exists(LOGO_PATH):
        try:
            slide.shapes.add_picture(LOGO_PATH, Inches(6.0), Inches(0.65), Inches(1.35), Inches(1.35))
        except Exception:
            pass

    # Institute Name
    inst_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.02), Inches(12.333), Inches(0.6))
    itf = inst_box.text_frame
    ip = itf.paragraphs[0]
    ip.text = "SAGAR COACHING CENTRE, BHAGWANPUR"
    ip.alignment = PP_ALIGN.CENTER
    ip.font.name = "Segoe UI"
    ip.font.size = Pt(16)
    ip.font.bold = True
    ip.font.color.rgb = C_GOLD

    # Badge Card
    badge_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.2), Inches(2.65), Inches(6.933), Inches(0.5))
    badge_card.fill.solid()
    badge_card.fill.fore_color.rgb = C_INDIGO
    badge_card.line.color.rgb = C_GOLD
    badge_card.line.width = Pt(1.5)
    btf = badge_card.text_frame
    bp = btf.paragraphs[0]
    bp.text = f"🔥 {badge_text} • {set_label} 🔥"
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Segoe UI"
    bp.font.size = Pt(13)
    bp.font.bold = True
    bp.font.color.rgb = C_WHITE

    # Main Title Box
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(12.333), Inches(1.2))
    ttf = title_box.text_frame
    ttf.word_wrap = True
    tp = ttf.paragraphs[0]
    tp.text = main_title
    tp.alignment = PP_ALIGN.CENTER
    tp.font.name = "Segoe UI"
    tp.font.size = Pt(28)
    tp.font.bold = True
    tp.font.color.rgb = C_WHITE

    tp_sub = ttf.add_paragraph()
    tp_sub.text = sub_desc
    tp_sub.alignment = PP_ALIGN.CENTER
    tp_sub.font.name = "Segoe UI"
    tp_sub.font.size = Pt(14)
    tp_sub.font.color.rgb = RGBColor(203, 213, 225)
    tp_sub.space_before = Pt(4)

    # Topics Card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.5), Inches(10.333), Inches(2.1))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(30, 41, 59)
    card.line.color.rgb = C_INDIGO
    card.line.width = Pt(1.5)
    ctf = card.text_frame
    ctf.word_wrap = True
    ctf.margin_top = Inches(0.12)
    cp0 = ctf.paragraphs[0]
    cp0.text = "🎯 इस अभ्यास सेट की मुख्य विशेषताएं (Live Class Features):"
    cp0.font.name = "Segoe UI"
    cp0.font.size = Pt(13.5)
    cp0.font.bold = True
    cp0.font.color.rgb = C_GOLD

    for t in topics_list:
        cp = ctf.add_paragraph()
        cp.text = f"✓ {t}"
        cp.font.name = "Segoe UI"
        cp.font.size = Pt(11.5)
        cp.font.color.rgb = C_WHITE
        cp.space_before = Pt(2)

    # Bottom Footer on Title
    bft = slide.shapes.add_textbox(Inches(0.5), Inches(6.75), Inches(12.333), Inches(0.5))
    bft_tf = bft.text_frame
    bp = bft_tf.paragraphs[0]
    bp.text = "मार्गदर्शक: सागर सर (NMMS Expert)  |  WhatsApp: +91 91101 13671  |  YouTube: Sagar Coaching Centre"
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Segoe UI"
    bp.font.size = Pt(11)
    bp.font.color.rgb = RGBColor(148, 163, 184)

def add_instructions_slide(prs, test_title, exam_type, total_slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, f"{test_title} — परीक्षा पैटर्न एवं अभ्यास निर्देश", 2, total_slides)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card.fill.solid()
    card.fill.fore_color.rgb = C_BLUE_CARD
    card.line.color.rgb = C_INDIGO
    card.line.width = Pt(2)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.3)

    p0 = tf.paragraphs[0]
    p0.text = f"📋 उत्तर प्रदेश NMMS {exam_type} परीक्षा योजना (Official Guidelines):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(18)
    p0.font.bold = True
    p0.font.color.rgb = C_NAVY

    if "MAT" in exam_type:
        rules = [
            "परीक्षा संरचना: मानसिक योग्यता परीक्षण (MAT) में कुल 90 बहुविकल्पीय प्रश्न (90 अंक) पूछे जाते हैं।",
            "इस अभ्यास सेट में UP NMMS के 10 वर्षों के ट्रेंड के आधार पर 50 अति-संभावित प्रश्न शामिल हैं।",
            "कवर किए गए प्रमुख विषय: श्रृंखला, सादृश्यता, वर्गीकरण, कोडिंग, लुप्त पद, गणितीय संक्रियाएं, रक्त संबंध, दिशा, वेन आरेख आदि।",
            "2-स्लाइड इंटरैक्टिव पद्धति: प्रश्न स्लाइड पर उत्तर पूरी तरह छुपा है ताकि विद्यार्थी पहले स्वयं विचार व अभ्यास करें।",
            "अगली स्लाइड पर सही विकल्प हरे रंग (Green Glow) में चमकेगा और नीचे आसान ट्रिक व व्याख्या प्रस्तुत होगी।",
            "लक्ष्य: MAT में 90 में से 85+ अंक प्राप्त करके मेरिट सूची में शीर्ष स्थान सुनिश्चित करना!"
        ]
    else:
        rules = [
            "परीक्षा संरचना: शैक्षिक योग्यता परीक्षण (SAT) में कुल 90 प्रश्न होते हैं (विज्ञान 35, सामाजिक विज्ञान 35, गणित 20)।",
            "इस 50-प्रश्नों के सेट में विषयवार विभाजन: विज्ञान = 22 प्रश्न | सामाजिक विज्ञान = 20 प्रश्न | गणित = 8 प्रश्न।",
            "पाठ्यक्रम: UP SCERT बेसिक शिक्षा परिषद् कक्षा 7 व 8 के संपूर्ण पाठ्यक्रम पर पूर्णतः आधारित।",
            "उत्तर प्रदेश विशेष: यूपी के ऐतिहासिक स्थल, 1857 क्रांति केंद्र, राष्ट्रीय उद्यान, नदियां व भूगोल के विशेष प्रश्न शामिल हैं।",
            "2-स्लाइड पद्धति: छात्र पहले प्रश्न देखकर विचार करेंगे, फिर अगली स्लाइड पर संपूर्ण व्याख्या व सूत्र समझेंगे।",
            "लक्ष्य: SAT में 90 में से 85+ अंक लाकर ₹48,000 की राष्ट्रीय छात्रवृत्ति पक्की करना!"
        ]

    for r in rules:
        p = tf.add_paragraph()
        p.text = f"•  {r}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(8)
        p.space_after = Pt(2)

def add_interactive_mcq_pair(prs, q_num, q_hi, q_en, options, correct_opt, explanation, exam_tag, subject_name, slide_num, total_slides):
    opt_labels = ["A", "B", "C", "D"]
    opt_pos = [
        (Inches(0.5), Inches(2.95), Inches(5.95), Inches(1.1)),
        (Inches(6.88), Inches(2.95), Inches(5.95), Inches(1.1)),
        (Inches(0.5), Inches(4.18), Inches(5.95), Inches(1.1)),
        (Inches(6.88), Inches(4.18), Inches(5.95), Inches(1.1)),
    ]

    # SLIDE 1: QUESTION ONLY (ANSWER HIDDEN)
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(s1, f"{subject_name} • प्र. #{q_num}  [{exam_tag}]", slide_num, total_slides)

    q_card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.15), Inches(12.333), Inches(1.68))
    q_card1.fill.solid()
    q_card1.fill.fore_color.rgb = C_BLUE_CARD
    q_card1.line.color.rgb = C_INDIGO
    q_card1.line.width = Pt(2)

    q_tf1 = q_card1.text_frame
    q_tf1.word_wrap = True
    q_tf1.margin_left = Inches(0.3)
    q_tf1.margin_top = Inches(0.1)
    q_tf1.margin_bottom = Inches(0.1)
    q_tf1.vertical_anchor = MSO_ANCHOR.MIDDLE

    qp1_hi = q_tf1.paragraphs[0]
    qp1_hi.text = f"प्र. {q_num}. {q_hi}"
    qp1_hi.alignment = PP_ALIGN.CENTER
    qp1_hi.font.name = "Segoe UI"
    qp1_hi.font.size = Pt(16.5)
    qp1_hi.font.bold = True
    qp1_hi.font.color.rgb = C_NAVY

    qp1_en = q_tf1.add_paragraph()
    qp1_en.text = f"Q. {q_num}. {q_en}"
    qp1_en.alignment = PP_ALIGN.CENTER
    qp1_en.font.name = "Segoe UI"
    qp1_en.font.size = Pt(14.5)
    qp1_en.font.bold = True
    qp1_en.font.color.rgb = RGBColor(37, 99, 235)
    qp1_en.space_before = Pt(6)

    for idx, (lbl, opt_text) in enumerate(zip(opt_labels, options)):
        left, top, w, h = opt_pos[idx]
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = RGBColor(203, 213, 225)
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = f"({lbl})  {opt_text}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_DARK

    # SLIDE 2: REVEAL SLIDE (ANSWER GLOWS GREEN + EXPLANATION)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(s2, f"{subject_name} • प्र. #{q_num}  [{exam_tag}]  —  [उत्तर व व्याख्या]", slide_num + 1, total_slides)

    q_card2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.15), Inches(12.333), Inches(1.68))
    q_card2.fill.solid()
    q_card2.fill.fore_color.rgb = C_BLUE_CARD
    q_card2.line.color.rgb = C_INDIGO
    q_card2.line.width = Pt(2)

    q_tf2 = q_card2.text_frame
    q_tf2.word_wrap = True
    q_tf2.margin_left = Inches(0.3)
    q_tf2.margin_top = Inches(0.1)
    q_tf2.margin_bottom = Inches(0.1)
    q_tf2.vertical_anchor = MSO_ANCHOR.MIDDLE

    qp2_hi = q_tf2.paragraphs[0]
    qp2_hi.text = f"प्र. {q_num}. {q_hi}"
    qp2_hi.alignment = PP_ALIGN.CENTER
    qp2_hi.font.name = "Segoe UI"
    qp2_hi.font.size = Pt(16.5)
    qp2_hi.font.bold = True
    qp2_hi.font.color.rgb = C_NAVY

    qp2_en = q_tf2.add_paragraph()
    qp2_en.text = f"Q. {q_num}. {q_en}"
    qp2_en.alignment = PP_ALIGN.CENTER
    qp2_en.font.name = "Segoe UI"
    qp2_en.font.size = Pt(14.5)
    qp2_en.font.bold = True
    qp2_en.font.color.rgb = RGBColor(37, 99, 235)
    qp2_en.space_before = Pt(6)

    for idx, (lbl, opt_text) in enumerate(zip(opt_labels, options)):
        left, top, w, h = opt_pos[idx]
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        is_correct = (lbl == correct_opt)

        if is_correct:
            card.fill.fore_color.rgb = RGBColor(220, 252, 231)  # Emerald light
            card.line.color.rgb = C_GREEN
            card.line.width = Pt(3)
        else:
            card.fill.fore_color.rgb = C_WHITE
            card.line.color.rgb = RGBColor(226, 232, 240)
            card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        prefix = "✅ " if is_correct else ""
        p.text = f"{prefix}({lbl})  {opt_text}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_GREEN if is_correct else RGBColor(148, 163, 184)

    # Bottom Explanation Card
    ans_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(5.42), Inches(12.333), Inches(1.5))
    ans_card.fill.solid()
    ans_card.fill.fore_color.rgb = C_AMBER_CARD
    ans_card.line.color.rgb = C_GOLD
    ans_card.line.width = Pt(1.5)

    atf = ans_card.text_frame
    atf.word_wrap = True
    atf.margin_left = Inches(0.3)
    atf.margin_top = Inches(0.12)

    ap0 = atf.paragraphs[0]
    ap0.text = f"✅ सही उत्तर / Correct Option: ({correct_opt})"
    ap0.font.name = "Segoe UI"
    ap0.font.size = Pt(15)
    ap0.font.bold = True
    ap0.font.color.rgb = C_GREEN

    ap1 = atf.add_paragraph()
    ap1.text = f"💡 व्याख्या एवं ट्रिक (Explanation): {explanation}"
    ap1.font.name = "Segoe UI"
    ap1.font.size = Pt(13)
    ap1.font.bold = False
    ap1.font.color.rgb = C_DARK
    ap1.space_before = Pt(3)

def add_outro_slide(prs, exam_title, total_slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, f"{exam_title} — सफलता का संकल्प एवं संपर्क सूत्र", total_slides, total_slides)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.333), Inches(5.2))
    card.fill.solid()
    card.fill.fore_color.rgb = C_NAVY
    card.line.color.rgb = C_GOLD
    card.line.width = Pt(2.5)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.4)

    p0 = tf.paragraphs[0]
    p0.text = "🏆 आपकी सफलता ही हमारा संकल्प है! 🏆"
    p0.alignment = PP_ALIGN.CENTER
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(22)
    p0.font.bold = True
    p0.font.color.rgb = C_GOLD

    p1 = tf.add_paragraph()
    p1.text = "SAGAR COACHING CENTRE, BHAGWANPUR"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    p1.space_before = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = "उत्तर प्रदेश NMMS छात्रवृत्ति परीक्षा (₹48,000) में चयनित होने वाले सभी भावी मेधावियों को अग्रिम शुभकामनाएं!"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(226, 232, 240)
    p2.space_before = Pt(10)

    boxes = [
        ("📱 WhatsApp हेल्पलाइन", "+91 91101 13671 (सागर सर)"),
        ("🌐 आधिकारिक वेबसाइट", "sagarcoaching.tech"),
        ("🎥 YouTube चैनल", "Sagar Coaching Centre (Free Classes)"),
        ("📍 पता / Centre Address", "भगवानपुर, सुपौल, बिहार - 852109")
    ]
    for b_title, b_val in boxes:
        p = tf.add_paragraph()
        p.text = f"•  {b_title}: {b_val}"
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.color.rgb = C_GOLD if "WhatsApp" in b_title else C_WHITE
        p.space_before = Pt(5)

# ----------------- MAIN BUILD LOGIC -----------------

def build_5_mat_sets():
    print("--- BUILDING 5 MAT SETS ---")
    os.makedirs(MAT_DIR, exist_ok=True)
    all_mat = m1.UP_MAT_QUESTIONS + m2.UP_MAT_ADDITIONAL + p3.UP_MAT_POOL3 + p4.UP_MAT_POOL4
    print(f"Total MAT Pool Available: {len(all_mat)}")

    topics_list = [
        "श्रृंखला परीक्षण (Number, Letter & Mixed Series)",
        "सादृश्यता एवं वर्गीकरण (Analogy & Classification)",
        "कोडिंग-डिकोडिंग, गणितीय संक्रियाएं एवं लुप्त पद",
        "रक्त संबंध, दिशा-दूरी, वेन आरेख एवं दर्पण प्रतिबिम्ब",
        "प्रत्येक प्रश्न के साथ सटीक ट्रिक एवं 2-स्लाइड इंटरैक्टिव व्याख्या"
    ]

    for set_idx in range(1, 6):
        set_num_str = f"{set_idx:02d}"
        set_label = f"अभ्यास सेट #{set_num_str} (Practice Set {set_idx})"
        ppt_name = f"UP_NMMS_MAT_Set_{set_num_str}.pptx"
        ppt_path = os.path.join(MAT_DIR, ppt_name)

        start_q = (set_idx - 1) * TOTAL_QUESTIONS_PER_SET
        end_q = start_q + TOTAL_QUESTIONS_PER_SET
        set_questions = all_mat[start_q:end_q]

        prs = create_base_presentation()

        # 1. Title Slide
        add_title_slide(
            prs,
            badge_text="UP NMMS 2024-25 महामंथन",
            main_title=f"मानसिक योग्यता परीक्षण (MAT) — अभ्यास सेट #{set_num_str}",
            sub_desc=f"परीक्षा में 100% सफलता हेतु 50 अति-संभावित प्रश्नों का महाभ्यास (प्र. 1 से 50)",
            topics_list=topics_list,
            set_label=f"सेट {set_num_str}"
        )

        # 2. Instructions Slide
        add_instructions_slide(
            prs,
            test_title=f"MAT अभ्यास सेट #{set_num_str}",
            exam_type="मानसिक योग्यता परीक्षण (MAT)",
            total_slides=TOTAL_SLIDES_PER_DECK
        )

        # 3. 50 Questions (2 slides each -> 100 slides)
        current_slide = 3
        for q_i, q in enumerate(set_questions, 1):
            add_interactive_mcq_pair(
                prs=prs,
                q_num=q_i,
                q_hi=q["question_hi"],
                q_en=q["question_en"],
                options=q["options"],
                correct_opt=q["correct"],
                explanation=q["explanation"],
                exam_tag=f"MAT सेट {set_num_str}",
                subject_name=f"मानसिक योग्यता (MAT) • {q.get('topic', 'तर्कशक्ति')}",
                slide_num=current_slide,
                total_slides=TOTAL_SLIDES_PER_DECK
            )
            current_slide += 2

        # 4. Outro Slide
        add_outro_slide(prs, f"UP NMMS MAT अभ्यास सेट #{set_num_str}", TOTAL_SLIDES_PER_DECK)

        prs.save(ppt_path)
        print(f"Generated: {ppt_name} ({len(prs.slides)} slides) -> {ppt_path}")

def build_5_sat_sets():
    print("\n--- BUILDING 5 SAT SETS ---")
    os.makedirs(SAT_DIR, exist_ok=True)
    all_sat = s1.UP_SAT_QUESTIONS + s2.UP_SAT_ADDITIONAL + p3.UP_SAT_POOL3 + p4.UP_SAT_POOL4
    print(f"Total SAT Pool Available: {len(all_sat)}")

    # Group SAT by main subjects: Science, SST, Math
    science_qs = []
    sst_qs = []
    math_qs = []

    for q in all_sat:
        sub = q.get("subject", "").lower()
        if "विज्ञान" in sub or "physics" in sub or "chemistry" in sub or "biology" in sub:
            science_qs.append(q)
        elif "गणित" in sub or "math" in sub:
            math_qs.append(q)
        else:
            sst_qs.append(q)

    print(f"Science Qs: {len(science_qs)}, SST Qs: {len(sst_qs)}, Math Qs: {len(math_qs)}")

    topics_list = [
        "विज्ञान (भौतिकी, रसायन, जीव विज्ञान — 22 प्रश्न)",
        "सामाजिक विज्ञान (इतिहास, भूगोल, नागरिक शास्त्र — 20 प्रश्न)",
        "गणित (अंकगणित, बीजगणित, ज्यामिति — 8 प्रश्न)",
        "उत्तर प्रदेश विशेष (इतिहास, संस्कृति, नदियां व भूगोल)",
        "प्रत्येक प्रश्न के साथ विस्तृत व्याख्या व 2-स्लाइड इंटरैक्टिव पद्धति"
    ]

    for set_idx in range(1, 6):
        set_num_str = f"{set_idx:02d}"
        set_label = f"अभ्यास सेट #{set_num_str} (Practice Set {set_idx})"
        ppt_name = f"UP_NMMS_SAT_Set_{set_num_str}.pptx"
        ppt_path = os.path.join(SAT_DIR, ppt_name)

        # 22 Science, 20 SST, 8 Math = 50 Questions per set
        sci_start = (set_idx - 1) * 22
        sci_set = science_qs[sci_start : sci_start + 22]

        sst_start = (set_idx - 1) * 20
        sst_set = sst_qs[sst_start : sst_start + 20]

        math_start = (set_idx - 1) * 8
        math_set = math_qs[math_start : math_start + 8]

        set_questions = sci_set + sst_set + math_set
        assert len(set_questions) == 50, f"Expected 50 Qs, got {len(set_questions)}"

        prs = create_base_presentation()

        # 1. Title Slide
        add_title_slide(
            prs,
            badge_text="UP NMMS 2024-25 महामंथन",
            main_title=f"शैक्षिक योग्यता परीक्षण (SAT) — अभ्यास सेट #{set_num_str}",
            sub_desc=f"विज्ञान (22), सामाजिक विज्ञान (20), गणित (8) • 50 अति-संभावित प्रश्न (प्र. 1 से 50)",
            topics_list=topics_list,
            set_label=f"सेट {set_num_str}"
        )

        # 2. Instructions Slide
        add_instructions_slide(
            prs,
            test_title=f"SAT अभ्यास सेट #{set_num_str}",
            exam_type="शैक्षिक योग्यता परीक्षण (SAT)",
            total_slides=TOTAL_SLIDES_PER_DECK
        )

        # 3. 50 Questions (2 slides each -> 100 slides)
        current_slide = 3
        for q_i, q in enumerate(set_questions, 1):
            add_interactive_mcq_pair(
                prs=prs,
                q_num=q_i,
                q_hi=q["question_hi"],
                q_en=q["question_en"],
                options=q["options"],
                correct_opt=q["correct"],
                explanation=q["explanation"],
                exam_tag=f"SAT सेट {set_num_str}",
                subject_name=f"{q.get('subject', 'शैक्षिक योग्यता')}",
                slide_num=current_slide,
                total_slides=TOTAL_SLIDES_PER_DECK
            )
            current_slide += 2

        # 4. Outro Slide
        add_outro_slide(prs, f"UP NMMS SAT अभ्यास सेट #{set_num_str}", TOTAL_SLIDES_PER_DECK)

        prs.save(ppt_path)
        print(f"Generated: {ppt_name} ({len(prs.slides)} slides) -> {ppt_path}")

if __name__ == "__main__":
    build_5_mat_sets()
    build_5_sat_sets()
    print("\nALL 10 UP NMMS PRACTICE SETS GENERATED SUCCESSFULLY!")
