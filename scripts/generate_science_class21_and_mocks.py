import sys
import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

from build_science_base import (
    MAHAMANTHAN_DIR,
    create_base_presentation,
    add_header_footer,
    add_title_slide,
    add_interactive_mcq_pair,
    add_memory_key_slide,
    finalize_and_save_presentation,
    C_NAVY, C_TEAL, C_GOLD, C_WHITE, C_DARK, C_MUTED, C_BLUE_CARD, C_AMBER_CARD
)

def add_instructions_slide(prs, title_text, bullet_points):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, title_text, 2, 105)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card.fill.solid()
    card.fill.fore_color.rgb = C_BLUE_CARD
    card.line.color.rgb = C_TEAL
    card.line.width = Pt(2)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.35)

    p0 = tf.paragraphs[0]
    p0.text = "📋 परीक्षा निर्देश एवं मूल्यांकन नियमावली (Exam Rules & Scoring Pattern):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(18)
    p0.font.bold = True
    p0.font.color.rgb = C_NAVY

    for pt in bullet_points:
        p = tf.add_paragraph()
        p.text = f"•  {pt}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(15)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(10)
        p.space_after = Pt(2)

def add_score_card_slide(prs, title_text):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, title_text, 104, 105)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = C_GOLD
    card.line.width = Pt(2)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.35)

    p0 = tf.paragraphs[0]
    p0.text = "🎯 स्व-मूल्यांकन एवं प्रदर्शन विश्लेषण (Self Evaluation Chart):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(19)
    p0.font.bold = True
    p0.font.color.rgb = C_NAVY

    levels = [
        "45+ सही उत्तर (उत्कृष्ट प्रदर्शन): State Merit List में सर्वोच्च स्थान सुनिश्चित! तैयारी पूर्ण है।",
        "38 - 44 सही उत्तर (सराहनीय प्रदर्शन): छात्रवृत्ति पक्की! कमजोर टॉपिक्स का 1 बार त्वरित रिवीजन करें।",
        "30 - 37 सही उत्तर (औसत प्रदर्शन): भौतिकी व रसायन के सूत्रों का पुनः अभ्यास करें।",
        "30 से कम सही उत्तर: क्लास 01 से 17 की थ्योरी स्लाइड्स को दोबारा पढ़ें और नोट्स बनाएं।"
    ]

    for lvl in levels:
        p = tf.add_paragraph()
        p.text = f"🏆  {lvl}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(15.5)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(12)

def add_institute_branding_slide(prs, badge_text, test_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, f"{test_title} — सफलता का संकल्प", 105, 105)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.6))
    card.fill.solid()
    card.fill.fore_color.rgb = C_NAVY
    card.line.color.rgb = C_GOLD
    card.line.width = Pt(2.5)

    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p0 = tf.paragraphs[0]
    p0.text = "SAGAR COACHING CENTRE, BHAGWANPUR"
    p0.alignment = PP_ALIGN.CENTER
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(24)
    p0.font.bold = True
    p0.font.color.rgb = C_GOLD

    p1 = tf.add_paragraph()
    p1.text = "निकट - माँ भगवती मंदिर, भगवानपुर (सुपौल, बिहार) • पिन - 852109"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(15)
    p1.font.color.rgb = RGBColor(203, 213, 225)
    p1.space_before = Pt(6)

    p2 = tf.add_paragraph()
    p2.text = f"🌟 {badge_text} : सम्पूर्ण पाठ्यक्रम सफलता किट 🌟"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = C_WHITE
    p2.space_before = Pt(16)

    p3 = tf.add_paragraph()
    p3.text = "मार्गदर्शक: सागर सर (NMMS King)  |  WhatsApp / Helpline: +91 91101 13671"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(16)
    p3.font.bold = True
    p3.font.color.rgb = C_GOLD
    p3.space_before = Pt(14)

    p4 = tf.add_paragraph()
    p4.text = "ऑनलाइन टेस्ट, क्लास नोट्स व PDF सामग्री हेतु वेबसाइट पर पधारें: sagarcoaching.tech"
    p4.alignment = PP_ALIGN.CENTER
    p4.font.name = "Segoe UI"
    p4.font.size = Pt(14)
    p4.font.color.rgb = RGBColor(226, 232, 240)
    p4.space_before = Pt(8)

def generate_50q_presentation(filename, class_badge, main_title, sub_desc, topics, questions_50, memory_sheet):
    prs = create_base_presentation()

    # 1. Title Slide
    add_title_slide(prs, class_badge, main_title, sub_desc, topics)

    # 2. Instructions Slide
    add_instructions_slide(
        prs,
        f"{main_title} — परीक्षा निर्देश",
        [
            "इस टेस्ट में राष्ट्रीय आय-सह-मेधा छात्रवृत्ति परीक्षा (NMMS) के विज्ञान (SAT) पाठ्यक्रम के कुल 50 प्रामाणिक प्रश्न शामिल हैं।",
            "प्रत्येक प्रश्न के लिए 1 अंक निर्धारित है। परीक्षा में कोई नकारात्मक अंकन (Negative Marking) नहीं है।",
            "प्रश्न स्लाइड पर उत्तर पूरी तरह छुपा हुआ है ताकि छात्र स्वयं विचार कर उत्तर दे सकें।",
            "अगली स्लाइड पर सही विकल्प हरे रंग में प्रकाशित होगा तथा नीचे संपूर्ण वैज्ञानिक तथ्य और व्याख्या दी गई है।",
            "लक्ष्य: विज्ञान के 35 प्रश्नों में से शत-प्रतिशत 35/35 अंक प्राप्त करना!"
        ]
    )

    # 3-102. 50 Questions in 2-Slide Pattern (100 slides)
    for q_idx, (q_num, question, opts, corr, expl, tag) in enumerate(questions_50):
        add_interactive_mcq_pair(
            prs,
            q_num,
            question,
            opts,
            corr,
            expl,
            tag,
            header_prefix="विज्ञान महामंथन PYQ"
        )

    # 103. Memory Keys Slide
    add_memory_key_slide(
        prs,
        f"{main_title} : त्वरित रिवीजन सूत्र",
        memory_sheet
    )

    # 104. Score Evaluation Slide
    add_score_card_slide(prs, f"{main_title} — स्कोर विश्लेषण")

    # 105. Institutional Slide
    add_institute_branding_slide(prs, class_badge, main_title)

    finalize_and_save_presentation(prs, filename, MAHAMANTHAN_DIR)


print("Science 50Q Engine ready.")
