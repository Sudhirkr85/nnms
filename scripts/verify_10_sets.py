# -*- coding: utf-8 -*-
import os
import sys
from pptx import Presentation

BASE_DIR = r"C:\Users\Orange\Downloads\New folder"
dirs = [
    os.path.join(BASE_DIR, "UP_NMMS_MAT_MahaManthan_5_Sets_SagarSir"),
    os.path.join(BASE_DIR, "UP_NMMS_SAT_MahaManthan_5_Sets_SagarSir")
]

forbidden = ["comment", "कमेंट", "10s", "timer", "thumbnail"]
found_issues = []

for d in dirs:
    for f in sorted(os.listdir(d)):
        if f.endswith(".pptx"):
            full_path = os.path.join(d, f)
            prs = Presentation(full_path)
            slide_count = len(prs.slides)
            if slide_count != 103:
                found_issues.append(f"{f}: Expected 103 slides, got {slide_count}")
            
            for s_idx, slide in enumerate(prs.slides, 1):
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        txt = shape.text_frame.text.lower()
                        for term in forbidden:
                            if term in txt:
                                found_issues.append(f"{f} Slide {s_idx}: Forbidden term '{term}'")

if found_issues:
    print("ISSUES FOUND:")
    for iss in found_issues:
        print(iss)
else:
    print("ALL 10 DECKS VERIFIED: Exactly 103 slides each, 0 forbidden terms found!")
