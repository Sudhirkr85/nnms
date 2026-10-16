import os
import sys
import io
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Encoding setup safely
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = r"d:\Work\E-Learning"
BASE_SST_DIR = r"C:\Users\Orange\Downloads\New folder\SST_Complete_Course_SagarSir"
CIVICS_DIR = os.path.join(BASE_SST_DIR, "01_Civics_नागरिक_शास्त्र")
HISTORY_DIR = os.path.join(BASE_SST_DIR, "02_History_इतिहास")
GEO_DIR = os.path.join(BASE_SST_DIR, "03_Geography_भूगोल")
MOCK_DIR = os.path.join(BASE_SST_DIR, "04_Full_Mock_Tests_महा_अभ्यास")

for d in [CIVICS_DIR, HISTORY_DIR, GEO_DIR, MOCK_DIR]:
    os.makedirs(d, exist_ok=True)

LOGO_PATH = os.path.join(ROOT_DIR, "public", "logo-circle-transparent.png")

# Colors
C_NAVY = RGBColor(30, 58, 138)       # #1E3A8A Primary Dark
C_GOLD = RGBColor(217, 119, 6)       # #D97706 Secondary Accent
C_LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC Soft BG
C_WHITE = RGBColor(255, 255, 255)
C_DARK = RGBColor(30, 41, 59)         # #1E293B Body Text
C_MUTED = RGBColor(100, 116, 139)    # #64748B Secondary Text
C_BLUE_CARD = RGBColor(239, 246, 255)# #EFF6FF Blue Light Card
C_AMBER_CARD = RGBColor(254, 243, 199)# #FEF3C7 Amber Light Card
C_GREEN = RGBColor(22, 101, 52)      # #166534 Success Answer
C_RED = RGBColor(185, 28, 28)        # #B91C1C Warning/Highlight

def create_base_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_header_footer(slide, title_text, slide_num, total_slides):
    # Header Background Card
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.2), Inches(12.533), Inches(0.85)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = C_NAVY
    header.line.color.rgb = C_GOLD
    header.line.width = Pt(1.5)
    
    # Logo in header (width 0.75 in, left 0.55 in -> ends at 1.30 in)
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(0.55), Inches(0.25), Inches(0.75), Inches(0.75))
        
    # Header Text in separate textbox (starts at 1.45 in, completely clear of logo!)
    tb = slide.shapes.add_textbox(Inches(1.45), Inches(0.2), Inches(11.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0)
    tf.margin_top = Inches(0.04)
    tf.margin_right = Inches(0.1)
    tf.margin_bottom = Inches(0.04)

    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Segoe UI"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.LEFT
    
    p_sub = tf.add_paragraph()
    p_sub.text = "सागर कोचिंग सेंटर भगवानपुर (सुपौल) | मार्गदर्शक: श्रवण कुमार 'सागर' (Sagar Sir) | YouTube: NMMS King Sagar Sir"
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = RGBColor(253, 230, 138) # Light gold
    
    # Footer
    footer = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(6.9), Inches(12.533), Inches(0.45)
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = C_LIGHT_BG
    footer.line.color.rgb = RGBColor(226, 232, 240)
    
    tf_f = footer.text_frame
    tf_f.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_f = tf_f.paragraphs[0]
    p_f.text = f"WhatsApp / Helpline: +91 91101 13671  |  Official Website: sagarcoaching.tech  |  Slide {slide_num} of {total_slides}"
    p_f.font.name = "Segoe UI"
    p_f.font.size = Pt(11)
    p_f.font.bold = True
    p_f.font.color.rgb = C_NAVY
    p_f.alignment = PP_ALIGN.CENTER

def add_title_slide(prs, class_badge, main_title, sub_desc, topics_list):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Big background card
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.4), Inches(12.333), Inches(6.7))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_NAVY
    bg.line.color.rgb = C_GOLD
    bg.line.width = Pt(3)
    
    # Logo
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(5.9), Inches(0.7), Inches(1.5), Inches(1.5))
        
    # Title Text Frame
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(11.333), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = "SAGAR COACHING CENTRE BHAGWANPUR"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(18)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(253, 230, 138)
    p0.alignment = PP_ALIGN.CENTER
    
    p1 = tf.add_paragraph()
    p1.text = f"NMMS All India (राष्ट्रीय आय-सह-मेधा छात्रवृत्ति)  •  {class_badge}"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(147, 197, 253)
    p1.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = main_title
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = C_WHITE
    p2.alignment = PP_ALIGN.CENTER
    
    p3 = tf.add_paragraph()
    p3.text = sub_desc
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(241, 245, 249)
    p3.alignment = PP_ALIGN.CENTER
    
    # Topic Highlights Card inside title slide
    hl_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.6), Inches(10.333), Inches(1.3))
    hl_card.fill.solid()
    hl_card.fill.fore_color.rgb = RGBColor(40, 75, 160)
    hl_card.line.color.rgb = C_GOLD
    hl_card.line.width = Pt(1.5)
    
    tf_hl = hl_card.text_frame
    tf_hl.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_h0 = tf_hl.paragraphs[0]
    p_h0.text = "🎯 आज की क्लास में क्या सीखेंगे (Key Topics):"
    p_h0.font.name = "Segoe UI"
    p_h0.font.size = Pt(14)
    p_h0.font.bold = True
    p_h0.font.color.rgb = RGBColor(253, 230, 138)
    p_h0.alignment = PP_ALIGN.CENTER
    
    p_h1 = tf_hl.add_paragraph()
    p_h1.text = "  ★  ".join(topics_list)
    p_h1.font.name = "Segoe UI"
    p_h1.font.size = Pt(13)
    p_h1.font.color.rgb = C_WHITE
    p_h1.alignment = PP_ALIGN.CENTER

    # Mentor Badge bottom
    m_box = slide.shapes.add_textbox(Inches(1.5), Inches(6.05), Inches(10.333), Inches(0.8))
    tf_m = m_box.text_frame
    p_m = tf_m.paragraphs[0]
    p_m.text = "मार्गदर्शक: श्रवण कुमार 'सागर' (Sagar Sir)  |  Helpline: +91 91101 13671  |  sagarcoaching.tech"
    p_m.font.name = "Segoe UI"
    p_m.font.size = Pt(13)
    p_m.font.bold = True
    p_m.font.color.rgb = RGBColor(226, 232, 240)
    p_m.alignment = PP_ALIGN.CENTER

def add_theory_slide(prs, title, points, vocab_dict=None, slide_num=1, total_slides=1, speaker_note=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, title, slide_num, total_slides)
    
    # Full Width Theory Card (12.333 inches)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.5))
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = RGBColor(203, 213, 225)
    card.line.width = Pt(1.5)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.4)
    tf.margin_right = Inches(0.5)
    
    p0 = tf.paragraphs[0]
    p0.text = "📖 मुख्य तथ्य एवं सिद्धांत (Key Concepts & Exam Notes):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(19)
    p0.font.bold = True
    p0.font.color.rgb = C_NAVY
    
    for pt in points:
        p = tf.add_paragraph()
        p.text = f"•  {pt}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(15.5)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(10)
        p.space_after = Pt(2)
        
    if speaker_note and slide.notes_slide:
        slide.notes_slide.notes_text_frame.text = speaker_note

def add_interactive_mcq_pair(prs, q_num, question, options, correct_opt, explanation, exam_tag, slide_num=1, total_slides=1, speaker_note="", header_prefix="PYQ अभ्यास प्रश्न"):
    opt_labels = ["A", "B", "C", "D"]
    opt_pos = [
        (Inches(0.5), Inches(2.9), Inches(5.95), Inches(1.1)),
        (Inches(6.88), Inches(2.9), Inches(5.95), Inches(1.1)),
        (Inches(0.5), Inches(4.15), Inches(5.95), Inches(1.1)),
        (Inches(6.88), Inches(4.15), Inches(5.95), Inches(1.1)),
    ]

    # ==========================================
    # SLIDE 1: QUESTION ONLY (ANSWER HIDDEN)
    # ==========================================
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(s1, f"{header_prefix} #{q_num}  [{exam_tag}]", slide_num, total_slides)

    # Question Card
    q_card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(1.5))
    q_card1.fill.solid()
    q_card1.fill.fore_color.rgb = C_BLUE_CARD
    q_card1.line.color.rgb = RGBColor(99, 102, 241)
    q_card1.line.width = Pt(2)

    q_tf1 = q_card1.text_frame
    q_tf1.word_wrap = True
    q_tf1.margin_left = Inches(0.3)
    q_tf1.vertical_anchor = MSO_ANCHOR.MIDDLE
    qp1 = q_tf1.paragraphs[0]
    qp1.text = f"प्र. {q_num}. {question}"
    qp1.font.name = "Segoe UI"
    qp1.font.size = Pt(18)
    qp1.font.bold = True
    qp1.font.color.rgb = C_NAVY

    # 4 Neutral Option Cards (Slide 1: Clean, no prompt card / no extra words)
    opt_pos = [
        (Inches(0.5), Inches(2.9), Inches(5.95), Inches(1.2)),
        (Inches(6.88), Inches(2.9), Inches(5.95), Inches(1.2)),
        (Inches(0.5), Inches(4.25), Inches(5.95), Inches(1.2)),
        (Inches(6.88), Inches(4.25), Inches(5.95), Inches(1.2)),
    ]

    for i, opt_text in enumerate(options):
        pos = opt_pos[i]
        lbl = opt_labels[i]
        opt_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pos[0], pos[1], pos[2], pos[3])
        opt_box.fill.solid()
        opt_box.fill.fore_color.rgb = C_WHITE
        opt_box.line.color.rgb = RGBColor(203, 213, 225)
        opt_box.line.width = Pt(1.5)

        otf = opt_box.text_frame
        otf.word_wrap = True
        otf.vertical_anchor = MSO_ANCHOR.MIDDLE
        otf.margin_left = Inches(0.2)
        op = otf.paragraphs[0]
        op.text = f"({lbl})  {opt_text}"
        op.font.name = "Segoe UI"
        op.font.size = Pt(16)
        op.font.bold = True
        op.font.color.rgb = C_DARK

    # ==========================================
    # SLIDE 2: ANSWER REVEALED (CORRECT GLOWS GREEN + EXPLANATION)
    # ==========================================
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(s2, f"{header_prefix} #{q_num}  [{exam_tag}]  —  [उत्तर व व्याख्या]", slide_num + 1, total_slides)

    # Same Question Card
    q_card2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(1.5))
    q_card2.fill.solid()
    q_card2.fill.fore_color.rgb = C_BLUE_CARD
    q_card2.line.color.rgb = RGBColor(99, 102, 241)
    q_card2.line.width = Pt(2)

    q_tf2 = q_card2.text_frame
    q_tf2.word_wrap = True
    q_tf2.margin_left = Inches(0.3)
    q_tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
    qp2 = q_tf2.paragraphs[0]
    qp2.text = f"प्र. {q_num}. {question}"
    qp2.font.name = "Segoe UI"
    qp2.font.size = Pt(18)
    qp2.font.bold = True
    qp2.font.color.rgb = C_NAVY

    # 4 Option Cards (Correct one glows GREEN with checkmark!)
    for i, opt_text in enumerate(options):
        pos = opt_pos[i]
        lbl = opt_labels[i]
        is_correct = (lbl == correct_opt)

        opt_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pos[0], pos[1], pos[2], pos[3])
        opt_box.fill.solid()

        if is_correct:
            opt_box.fill.fore_color.rgb = C_GREEN
            opt_box.line.color.rgb = C_GOLD
            opt_box.line.width = Pt(3)
        else:
            opt_box.fill.fore_color.rgb = C_WHITE
            opt_box.line.color.rgb = RGBColor(226, 232, 240)
            opt_box.line.width = Pt(1)

        otf = opt_box.text_frame
        otf.word_wrap = True
        otf.vertical_anchor = MSO_ANCHOR.MIDDLE
        otf.margin_left = Inches(0.2)
        op = otf.paragraphs[0]

        if is_correct:
            op.text = f"✅ ({lbl})  {opt_text}"
            op.font.name = "Segoe UI"
            op.font.size = Pt(16)
            op.font.bold = True
            op.font.color.rgb = C_WHITE
        else:
            op.text = f"({lbl})  {opt_text}"
            op.font.name = "Segoe UI"
            op.font.size = Pt(15)
            op.font.color.rgb = RGBColor(148, 163, 184) # Muted

    # Answer and Explanation Banner at Bottom
    exp_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(5.6), Inches(12.333), Inches(1.2))
    exp_card.fill.solid()
    exp_card.fill.fore_color.rgb = RGBColor(240, 253, 244) # Soft green
    exp_card.line.color.rgb = C_GREEN
    exp_card.line.width = Pt(2)

    etf = exp_card.text_frame
    etf.word_wrap = True
    etf.margin_left = Inches(0.25)
    etf.vertical_anchor = MSO_ANCHOR.MIDDLE

    ep0 = etf.paragraphs[0]
    ep0.text = f"✅ सही उत्तर: ({correct_opt})"
    ep0.font.name = "Segoe UI"
    ep0.font.size = Pt(16)
    ep0.font.bold = True
    ep0.font.color.rgb = C_GREEN

    ep1 = etf.add_paragraph()
    ep1.text = f"💡 व्याख्या: {explanation}"
    ep1.font.name = "Segoe UI"
    ep1.font.size = Pt(13.5)
    ep1.font.color.rgb = RGBColor(20, 83, 45)

    if speaker_note and s2.notes_slide:
        s2.notes_slide.notes_text_frame.text = speaker_note

def add_mcq_slide(prs, q_num, question, options, correct_opt, explanation, exam_tag, slide_num=1, total_slides=1, speaker_note="", header_prefix="PYQ अभ्यास प्रश्न"):
    # Backward compatibility: delegates to pair
    add_interactive_mcq_pair(prs, q_num, question, options, correct_opt, explanation, exam_tag, slide_num, total_slides, speaker_note, header_prefix)

def finalize_and_save_presentation(prs, filename, target_dir):
    """Dynamically fixes all slide numbers in footers and safely writes PPTX"""
    total_slides = len(prs.slides)
    for idx, slide in enumerate(prs.slides):
        current_num = idx + 1
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame:
                txt = shape.text_frame.text
                if "sagarcoaching.tech" in txt and "Slide " in txt:
                    p = shape.text_frame.paragraphs[0]
                    p.text = f"WhatsApp / Helpline: +91 91101 13671  |  Official Website: sagarcoaching.tech  |  Slide {current_num} of {total_slides}"
                    p.font.name = "Segoe UI"
                    p.font.size = Pt(11)
                    p.font.bold = True
                    p.font.color.rgb = C_NAVY
                    p.alignment = PP_ALIGN.CENTER

    os.makedirs(target_dir, exist_ok=True)
    out_path = os.path.join(target_dir, filename)
    try:
        prs.save(out_path)
        print(f"--> [SAVED] {filename} ({total_slides} slides)")
        return out_path
    except PermissionError:
        alt_name = filename.replace(".pptx", "_Updated.pptx")
        alt_path = os.path.join(target_dir, alt_name)
        prs.save(alt_path)
        print(f"--> [LOCKED - Saved as {alt_name}] ({total_slides} slides)")
        return alt_path

def add_memory_key_slide(prs, title, key_items, slide_num, total_slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, title, slide_num, total_slides)
    
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.333), Inches(5.5))
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = C_GOLD
    card.line.width = Pt(2.5)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.4)
    
    p0 = tf.paragraphs[0]
    p0.text = "🔑 1-मिनट क्विक मेमोरी सूत्र (Revision Formulae):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(20)
    p0.font.bold = True
    p0.font.color.rgb = C_GOLD
    
    for item in key_items:
        p = tf.add_paragraph()
        p.text = f"👉  {item}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_NAVY
        p.space_before = Pt(10)

def add_homework_slide(prs, hw_questions, slide_num, total_slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_footer(slide, "आज का अभ्यास कार्य (Homework Questions)", slide_num, total_slides)
    
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.733), Inches(5.3))
    card.fill.solid()
    card.fill.fore_color.rgb = C_NAVY
    card.line.color.rgb = C_GOLD
    card.line.width = Pt(2.5)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.4)
    
    p0 = tf.paragraphs[0]
    p0.text = "📝 आज का अभ्यास कार्य (Homework):"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(22)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(253, 230, 138)
    
    for idx, hw in enumerate(hw_questions):
        p = tf.add_paragraph()
        p.text = f"प्रश्न {idx+1}: {hw}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(14)

print("Base builder functions defined successfully.")
